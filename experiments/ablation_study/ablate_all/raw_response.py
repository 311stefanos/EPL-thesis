# --- receipt_agent.py --- #
''' Imports '''

import json
import os
import sys
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

from dotenv import load_dotenv
from langchain_core.messages import AIMessage, BaseMessage, HumanMessage, SystemMessage
from langchain_core.tools import tool
from langgraph.graph import END, START, MessagesState, StateGraph
from langgraph.prebuilt import ToolNode

from utils.utils import myChatOpenAI, safe_invoke, will_tool_call

from creations.receipt_agent import receipt_agent_prompts as prompts
from creations.receipt_agent.receipt_processor import receipt_processor_app
from creations.receipt_agent.receipt_storage import (
    filter_records_by_date,
    load_records,
    parse_items,
)

''' Constants '''

load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent.parent / '.env')

DEBUG = os.getenv('DEBUG', 'false').strip().lower() in ('1', 'true', 'yes')

''' Schemas '''

class ReceiptAgentSchema(MessagesState):
    # Path to the receipt photo the user sent; empty/absent when the input is a text question.
    image_path: Optional[str]
    # The user's spending question, extracted from the latest message.
    question: Optional[str]
    # Final natural-language answer produced for a spending question.
    answer: Optional[str]
    # Path of the Excel file receipts are stored in.
    excel_path: Optional[str]
    # High-level outcome of the run: 'receipt_saved' | 'question_answered' | 'error'.
    status: Optional[str]
    # Operational error description; empty when everything succeeded.
    error: Optional[str]

''' Tools '''

@tool
def get_all_records() -> str:
    """Return every receipt record currently stored in the Excel file.

    Use this tool when the user's question needs the full detail of past
    purchases, such as questions about specific stores, locations, dates or
    individual receipts, or when the aggregated tools do not fit the question.

    Args:
        None.

    Returns:
        A JSON string shaped as {"count": <int>, "records": [{"cost": number,
        "date": "YYYY-MM-DD", "items": "item (quantity x price CURRENCY), ...",
        "location": str, "category": str}, ...]}.
    """
    records = load_records()
    return json.dumps({'count': len(records), 'records': records}, ensure_ascii=False)


@tool
def get_category_totals(start_date: str = '', end_date: str = '') -> str:
    """Return total spending per category, optionally within an inclusive date range.

    Use this tool for questions like "how much did I spend on groceries this
    month?" — build the ISO date range from today's date first.

    Args:
        start_date: Optional inclusive range start as ISO "YYYY-MM-DD". Empty means no lower bound.
        end_date: Optional inclusive range end as ISO "YYYY-MM-DD". Empty means no upper bound.

    Returns:
        A JSON string shaped as {"start_date": str, "end_date": str,
        "categories": [{"category": str, "total_cost": number, "receipt_count": int,
        "currencies": [str]}, ...]} sorted from highest to lowest total cost.
    """
    records = filter_records_by_date(load_records(), start_date, end_date)
    totals: Dict[str, Dict[str, Any]] = {}
    for record in records:
        category = record.get('category') or 'Uncategorized'
        entry = totals.setdefault(
            category,
            {'category': category, 'total_cost': 0.0, 'receipt_count': 0, 'currencies': set()},
        )
        entry['total_cost'] += float(record.get('cost') or 0.0)
        entry['receipt_count'] += 1
        for item in parse_items(record.get('items', '')):
            if item.get('currency'):
                entry['currencies'].add(item['currency'])
    ranked = sorted(totals.values(), key=lambda entry: entry['total_cost'], reverse=True)
    for entry in ranked:
        entry['currencies'] = sorted(entry['currencies'])
    return json.dumps(
        {'start_date': start_date, 'end_date': end_date, 'categories': ranked},
        ensure_ascii=False,
    )


@tool
def get_monthly_totals(category: str = '') -> str:
    """Return total spending per calendar month, optionally filtered to one category.

    Use this tool for trend questions like "how much did I spend per month on
    dining?" or "what were my biggest spending months?".

    Args:
        category: Optional category name (case-insensitive) to filter by. Empty means all categories.

    Returns:
        A JSON string shaped as {"category": str, "months": [{"month": "YYYY-MM",
        "total_cost": number, "receipt_count": int, "currencies": [str]}, ...]}
        sorted from earliest to latest month.
    """
    records = load_records()
    wanted = category.strip().lower()
    totals: Dict[str, Dict[str, Any]] = {}
    for record in records:
        if wanted and (record.get('category') or '').strip().lower() != wanted:
            continue
        month = (record.get('date') or '')[:7]
        if not month:
            continue
        entry = totals.setdefault(
            month, {'month': month, 'total_cost': 0.0, 'receipt_count': 0, 'currencies': set()}
        )
        entry['total_cost'] += float(record.get('cost') or 0.0)
        entry['receipt_count'] += 1
        for item in parse_items(record.get('items', '')):
            if item.get('currency'):
                entry['currencies'].add(item['currency'])
    months = [totals[key] for key in sorted(totals)]
    for entry in months:
        entry['currencies'] = sorted(entry['currencies'])
    return json.dumps({'category': category, 'months': months}, ensure_ascii=False)


@tool
def get_top_items(limit: int = 3, start_date: str = '', end_date: str = '') -> str:
    """Return the items the user spent the most money on, ranked by total spend.

    Each recorded line item contributes (quantity x unit price) to its item
    total, and identical item names are merged across receipts. Use this tool
    for questions like "what are the top 3 items I spent the most on?".

    Args:
        limit: How many top items to return (minimum 1, default 3).
        start_date: Optional inclusive range start as ISO "YYYY-MM-DD".
        end_date: Optional inclusive range end as ISO "YYYY-MM-DD".

    Returns:
        A JSON string shaped as {"limit": int, "top_items": [{"name": str,
        "total_spend": number, "total_quantity": number, "currencies": [str]}, ...]}
        sorted from highest to lowest spend.
    """
    records = filter_records_by_date(load_records(), start_date, end_date)
    totals: Dict[str, Dict[str, Any]] = {}
    for record in records:
        for item in parse_items(record.get('items', '')):
            key = item['name'].strip().lower()
            if not key:
                continue
            entry = totals.setdefault(
                key,
                {
                    'name': item['name'].strip(),
                    'total_spend': 0.0,
                    'total_quantity': 0.0,
                    'currencies': set(),
                },
            )
            entry['total_spend'] += item['quantity'] * item['unit_price']
            entry['total_quantity'] += item['quantity']
            if item.get('currency'):
                entry['currencies'].add(item['currency'])
    safe_limit = max(1, int(limit or 3))
    ranked = sorted(totals.values(), key=lambda entry: entry['total_spend'], reverse=True)[:safe_limit]
    for entry in ranked:
        entry['currencies'] = sorted(entry['currencies'])
    return json.dumps({'limit': safe_limit, 'top_items': ranked}, ensure_ascii=False)


SPENDING_TOOLS = [get_all_records, get_category_totals, get_monthly_totals, get_top_items]

''' LLM '''

# Tool-bound model that decides which spending queries to run, then answers.
answer_llm = myChatOpenAI(temperature=0.0).bind_tools(SPENDING_TOOLS)

''' Helpful Functions '''

def _message_text(message: BaseMessage) -> str:
    """Extract the plain-text portion of a message, joining multimodal parts."""
    content = message.content
    if isinstance(content, str):
        return content.strip()
    if isinstance(content, list):
        parts: List[str] = []
        for block in content:
            if isinstance(block, dict) and block.get('type') == 'text':
                parts.append(str(block.get('text', '')))
            elif isinstance(block, str):
                parts.append(block)
        return ' '.join(parts).strip()
    return ''


def _last_human_text(messages: List[BaseMessage]) -> str:
    """Return the text of the most recent human message, or '' when none exists."""
    for message in reversed(messages):
        if isinstance(message, HumanMessage):
            return _message_text(message)
    return ''

''' Nodes '''

def prepare_input_node(state: ReceiptAgentSchema) -> Dict[str, Any]:
    """Normalize the incoming request: keep photos as-is, extract question text."""
    updates: Dict[str, Any] = {'error': ''}
    if state.get('image_path'):
        return updates
    question = (state.get('question') or '').strip()
    if not question:
        question = _last_human_text(state.get('messages', []))
    if question:
        updates['question'] = question
    else:
        updates['error'] = 'No receipt photo or spending question was found in the request.'
    return updates


def process_receipt_node(state: ReceiptAgentSchema) -> Dict[str, Any]:
    """Run the receipt-processing subgraph on the supplied photo."""
    try:
        result = receipt_processor_app.invoke({'image_path': state.get('image_path') or ''})
    except Exception:
        if DEBUG:
            import traceback

            traceback.print_exc()
        return {'error': 'The receipt processing pipeline failed unexpectedly.', 'status': 'error'}
    return {
        'messages': result.get('messages', []),
        'excel_path': result.get('excel_path', ''),
        'status': result.get('status', ''),
        'error': result.get('error', '') or '',
    }


def answer_question_node(state: ReceiptAgentSchema) -> Dict[str, Any]:
    """Ask the tool-bound LLM to query spending data or produce the final answer."""
    question = state.get('question', '')
    today = date.today().isoformat()
    prompt = prompts.ANSWER_QUESTION_PROMPT.format(question=question, today=today)
    messages = [SystemMessage(content=prompt)] + list(state.get('messages', []))
    try:
        response = safe_invoke(answer_llm, messages)
    except Exception:
        if DEBUG:
            import traceback

            traceback.print_exc()
        return {
            'messages': [AIMessage(content='⚠️ Sorry, I could not analyze your spending right now. Please try again in a moment.')],
            'error': 'The spending assistant failed to respond.',
        }
    return {'messages': [response]}


def finalize_answer_node(state: ReceiptAgentSchema) -> Dict[str, Any]:
    """Copy the final assistant message into the dedicated answer field."""
    answer = ''
    for message in reversed(state.get('messages', [])):
        if isinstance(message, AIMessage):
            answer = _message_text(message)
            break
    status = 'error' if state.get('error') else 'question_answered'
    return {'answer': answer, 'status': status}


def handle_error_node(state: ReceiptAgentSchema) -> Dict[str, Any]:
    """Produce the user-facing message for the failure recorded in state."""
    error = state.get('error') or 'Something went wrong while handling your request.'
    if DEBUG:
        print(f'[receipt_agent] run failed: {error}')
    return {
        'messages': [AIMessage(content=f'⚠️ Sorry, I could not process your request: {error}')],
        'status': 'error',
    }

''' Conditional Functions '''

def route_input(state: ReceiptAgentSchema) -> Literal['process_receipt', 'answer_question', 'handle_error']:
    """Send photos to the receipt pipeline and text to the spending assistant."""
    if state.get('image_path'):
        return 'process_receipt'
    if state.get('error'):
        return 'handle_error'
    return 'answer_question'


def route_after_receipt(state: ReceiptAgentSchema) -> Literal['handle_error', 'done']:
    """Fail visibly when the receipt pipeline reported a problem."""
    return 'handle_error' if state.get('error') else 'done'


def route_after_answer(state: ReceiptAgentSchema) -> Literal['spending_tools', 'finalize_answer']:
    """Loop through tool execution until the assistant produces its final answer."""
    if will_tool_call(state.get('messages', [])):
        return 'spending_tools'
    return 'finalize_answer'

''' Graph '''

receipt_agent_graph = StateGraph(ReceiptAgentSchema)

receipt_agent_graph.add_node('prepare_input', prepare_input_node)
receipt_agent_graph.add_node('process_receipt', process_receipt_node)
receipt_agent_graph.add_node('answer_question', answer_question_node)
receipt_agent_graph.add_node('spending_tools', ToolNode(SPENDING_TOOLS))
receipt_agent_graph.add_node('finalize_answer', finalize_answer_node)
receipt_agent_graph.add_node('handle_error', handle_error_node)

receipt_agent_graph.add_edge(START, 'prepare_input')
receipt_agent_graph.add_conditional_edges(
    'prepare_input',
    route_input,
    {
        'process_receipt': 'process_receipt',
        'answer_question': 'answer_question',
        'handle_error': 'handle_error',
    },
)
receipt_agent_graph.add_conditional_edges(
    'process_receipt',
    route_after_receipt,
    {'handle_error': 'handle_error', 'done': END},
)
receipt_agent_graph.add_conditional_edges(
    'answer_question',
    route_after_answer,
    {'spending_tools': 'spending_tools', 'finalize_answer': 'finalize_answer'},
)
receipt_agent_graph.add_edge('spending_tools', 'answer_question')
receipt_agent_graph.add_edge('finalize_answer', END)
receipt_agent_graph.add_edge('handle_error', END)

receipt_agent_app = receipt_agent_graph.compile()

''' Testing '''

if __name__ == '__main__':
    # Example 1: ask a spending question (reads the Excel file, no photo needed).
    demo_question = 'How much did I spend on groceries this month, and what are the top 3 items I spent the most on?'
    demo_result = receipt_agent_app.invoke(
        {'messages': [HumanMessage(content=demo_question)]},
        config={'recursion_limit': 40},
    )
    print('Status:', demo_result.get('status'))
    print('Answer:', demo_result.get('answer'))

    # Example 2: process a receipt photo passed as the first command-line argument.
    if len(sys.argv) > 1:
        photo_result = receipt_agent_app.invoke(
            {
                'image_path': sys.argv[1],
                'messages': [HumanMessage(content='Please save this receipt.')],
            },
            config={'recursion_limit': 40},
        )
        print('Status:', photo_result.get('status'))
        for message in photo_result.get('messages', []):
            if isinstance(message, AIMessage):
                print(message.content)

# --- receipt_agent_prompts.py --- #

ANSWER_QUESTION_PROMPT = """\\
# Role
You are a personal spending assistant with access to the user's receipt records.

# Objective
Answer the user's question about their recorded spending accurately and concisely, using the tools available to you.

# Inputs
- User question: {question}
- Today's date: {today}. Use it to resolve relative periods such as "this month", "last week", or "this year".

# Available Tools
- get_all_records(): returns every stored receipt (cost, date, items, location, category). Use it for questions about specific purchases, stores, dates, or any detail the aggregated tools cannot provide.
- get_category_totals(start_date, end_date): total cost per category within an optional inclusive ISO date range. Use it for "how much did I spend on <category>" questions.
- get_monthly_totals(category): total cost per calendar month, optionally filtered to one category. Use it for month-by-month trends.
- get_top_items(limit, start_date, end_date): the items with the highest total spend (quantity x unit price, merged across receipts). Use it for "top N items" questions.

# Instructions
1. Decide which tool(s) are needed to answer the question and call them before answering. Usually one or two calls are enough.
2. Translate relative periods into concrete ISO date ranges before calling tools (e.g. "this month" -> from the first day of the current month to today).
3. Base every number in your answer strictly on the data returned by the tools, and do the arithmetic carefully.
4. If the records use several currencies, report the totals per currency instead of mixing them.
5. If there are no records at all, or none matching the question, say so clearly instead of guessing.

# Hard Rules
- Never invent or estimate amounts that did not come from a tool result.
- Do not keep calling tools once you have enough data; always finish with a final answer.
- Answer in the language of the user's question.

# Output
A short, friendly, direct answer (2-5 sentences, or a small list when ranking items), with every amount accompanied by its currency.
"""

# --- receipt_processor.py --- #
''' Imports '''

import base64
import os
import sys
import traceback
from datetime import date
from pathlib import Path
from typing import Any, Dict, List, Literal, Optional

from dotenv import load_dotenv
from langchain_core.messages import AIMessage, HumanMessage, SystemMessage
from langgraph.graph import END, START, MessagesState, StateGraph
from pydantic import BaseModel, Field

from utils.utils import myChatOpenAI, safe_invoke

from creations.receipt_agent import receipt_processor_prompts as prompts
from creations.receipt_agent.receipt_storage import append_receipt_record, format_items_string

''' Constants '''

load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent.parent / '.env')

DEBUG = os.getenv('DEBUG', 'false').strip().lower() in ('1', 'true', 'yes')

# Vision-capable model called through OpenRouter to read receipt photos.
VISION_MODEL = os.getenv('OPENROUTER_VISION_MODEL', 'openai/gpt-4o-mini')

MIME_BY_SUFFIX = {
    '.jpg': 'image/jpeg',
    '.jpeg': 'image/jpeg',
    '.png': 'image/png',
    '.webp': 'image/webp',
}
SUPPORTED_IMAGE_SUFFIXES = set(MIME_BY_SUFFIX)

# Maximum accepted photo size (vision providers commonly cap around 20 MB).
MAX_IMAGE_BYTES = 20 * 1024 * 1024

ALLOWED_CATEGORIES: List[str] = [
    'Groceries',
    'Dining',
    'Transport',
    'Utilities',
    'Shopping',
    'Health',
    'Entertainment',
    'Travel',
    'Education',
    'Other',
]

''' Schemas '''

class ReceiptItem(BaseModel):
    """A single purchased line item as printed on the receipt."""
    name: str = Field(description='Clean item name without commas or parentheses.')
    quantity: float = Field(description='Number of units purchased for this line.')
    unit_price: float = Field(description='Price of a single unit in the receipt currency.')


class ReceiptData(BaseModel):
    """Structured data extracted from one receipt photo."""
    location: str = Field(description='Store or merchant name, with branch or city when visible.')
    date: str = Field(description='Receipt date in ISO format (YYYY-MM-DD).')
    currency: str = Field(description='ISO 4217 currency code (e.g. USD, EUR, ILS) or the symbol found on the receipt.')
    total_cost: float = Field(description='Final total amount paid as printed on the receipt.')
    items: List[ReceiptItem] = Field(description='All purchased line items with quantity and unit price.')


class ReceiptCategory(BaseModel):
    """The single spending category assigned to a receipt."""
    category: str = Field(description='Exactly one category from the allowed category list.')


class ReceiptProcessorSchema(MessagesState):
    # Path of the receipt photo being processed.
    image_path: Optional[str]
    # Base64 data URL of the photo, produced by the image loader.
    image_data_url: Optional[str]
    # Structured receipt payload extracted by the vision model (see ReceiptData).
    receipt_data: Optional[dict]
    # Spending category assigned to the receipt.
    category: Optional[str]
    # Path of the Excel file the receipt was appended to.
    excel_path: Optional[str]
    # High-level outcome: 'receipt_saved' | 'error'.
    status: Optional[str]
    # Operational error description; empty when everything succeeded.
    error: Optional[str]

''' LLM '''

# Vision model with structured output that reads the receipt photo.
extraction_llm = myChatOpenAI(temperature=0.0, model=VISION_MODEL).with_structured_output(ReceiptData)

# Text model with structured output that assigns the spending category.
category_llm = myChatOpenAI(temperature=0.0).with_structured_output(ReceiptCategory)

''' Helpful Functions '''

def _encode_image(image_path: str) -> str:
    """Read an image file and return it as a base64 data URL for the vision model.

    Raises ValueError with a user-safe message when the path is missing, the
    file does not exist, the type is unsupported, or the file is too large.
    """
    if not image_path:
        raise ValueError('No receipt photo path was provided.')
    path = Path(image_path).expanduser()
    if not path.is_file():
        raise ValueError(f'Receipt photo not found: {image_path}')
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_IMAGE_SUFFIXES:
        supported = ', '.join(sorted(SUPPORTED_IMAGE_SUFFIXES))
        raise ValueError(f'Unsupported photo type "{suffix}". Supported types: {supported}.')
    data = path.read_bytes()
    if len(data) > MAX_IMAGE_BYTES:
        raise ValueError('The receipt photo is too large (the limit is 20 MB).')
    encoded = base64.b64encode(data).decode('ascii')
    return f'data:{MIME_BY_SUFFIX[suffix]};base64,{encoded}'


def _normalize_category(raw_category: str) -> str:
    """Map a model-proposed category onto the allowed list, defaulting to Other."""
    candidate = (raw_category or '').strip().lower()
    for allowed in ALLOWED_CATEGORIES:
        if candidate == allowed.lower():
            return allowed
    return 'Other'

''' Nodes '''

def load_image_node(state: ReceiptProcessorSchema) -> Dict[str, Any]:
    """Validate the photo path and encode the image for the vision model."""
    try:
        data_url = _encode_image(state.get('image_path') or '')
    except ValueError as exc:
        return {'error': str(exc), 'status': 'error'}
    except OSError:
        if DEBUG:
            traceback.print_exc()
        return {'error': 'The receipt photo could not be read.', 'status': 'error'}
    return {'image_data_url': data_url, 'error': ''}


def extract_receipt_node(state: ReceiptProcessorSchema) -> Dict[str, Any]:
    """Read the receipt photo with the vision model and extract structured data."""
    data_url = state.get('image_data_url') or ''
    if not data_url:
        return {'error': 'No receipt photo was loaded, so nothing could be extracted.', 'status': 'error'}
    today = date.today().isoformat()
    year = today[:4]
    prompt = prompts.EXTRACT_RECEIPT_PROMPT.format(today=today, year=year)
    message = HumanMessage(
        content=[
            {'type': 'text', 'text': 'Extract the data from this receipt photo.'},
            {'type': 'image_url', 'image_url': {'url': data_url}},
        ]
    )
    try:
        response = safe_invoke(extraction_llm, [SystemMessage(content=prompt), message])
    except Exception:
        if DEBUG:
            traceback.print_exc()
        return {'error': 'Reading the receipt photo failed; no data could be extracted.', 'status': 'error'}
    if response is None:
        return {'error': 'Reading the receipt photo failed; no data could be extracted.', 'status': 'error'}
    return {'receipt_data': response.model_dump(), 'error': ''}


def categorize_receipt_node(state: ReceiptProcessorSchema) -> Dict[str, Any]:
    """Determine the spending category for the extracted receipt items."""
    receipt = state.get('receipt_data') or {}
    items_text = format_items_string(receipt)
    location = str(receipt.get('location') or 'Unknown')
    prompt = prompts.CATEGORIZE_RECEIPT_PROMPT.format(
        items=items_text or 'No item details were extracted.',
        location=location,
        categories=', '.join(ALLOWED_CATEGORIES),
    )
    raw_category = ''
    try:
        response = safe_invoke(
            category_llm,
            [SystemMessage(content=prompt), HumanMessage(content='Assign the spending category.')],
        )
        raw_category = getattr(response, 'category', '')
    except Exception:
        if DEBUG:
            traceback.print_exc()
    return {'category': _normalize_category(raw_category)}


def save_receipt_node(state: ReceiptProcessorSchema) -> Dict[str, Any]:
    """Append the categorized receipt as a new row in the Excel file."""
    receipt = state.get('receipt_data') or {}
    category = state.get('category') or 'Other'
    try:
        excel_path = append_receipt_record(receipt, category)
    except Exception as exc:
        if DEBUG:
            traceback.print_exc()
        return {'error': f'Failed to write the receipt to the Excel file: {exc}', 'status': 'error'}
    return {'excel_path': str(excel_path), 'status': 'receipt_saved', 'error': ''}


def compose_confirmation_node(state: ReceiptProcessorSchema) -> Dict[str, Any]:
    """Build the deterministic confirmation message shown to the user."""
    receipt = state.get('receipt_data') or {}
    category = state.get('category') or 'Other'
    total_cost = float(receipt.get('total_cost') or 0.0)
    currency = str(receipt.get('currency') or '').strip()
    lines = [
        '✅ Receipt saved!',
        f'• Location: {receipt.get("location") or "Unknown"}',
        f'• Date: {receipt.get("date") or "Unknown"}',
        f'• Total: {total_cost:.2f} {currency}'.strip(),
        f'• Category: {category}',
        f'• Items: {format_items_string(receipt) or "No item details extracted"}',
    ]
    return {'messages': [AIMessage(content='\
'.join(lines))]}


def handle_error_node(state: ReceiptProcessorSchema) -> Dict[str, Any]:
    """Mark the failure; the root graph renders the user-facing message."""
    if DEBUG and state.get('error'):
        print(f'[receipt_processor] error: {state["error"]}')
    return {'status': 'error'}

''' Conditional Functions '''

def route_after_image(state: ReceiptProcessorSchema) -> Literal['extract_receipt', 'handle_error']:
    """Skip extraction when the photo could not be loaded."""
    return 'handle_error' if state.get('error') else 'extract_receipt'


def route_after_save(state: ReceiptProcessorSchema) -> Literal['compose_confirmation', 'handle_error']:
    """Only confirm success when the row was actually written."""
    return 'handle_error' if state.get('error') else 'compose_confirmation'

''' Graph '''

receipt_processor_graph = StateGraph(ReceiptProcessorSchema)

receipt_processor_graph.add_node('load_image', load_image_node)
receipt_processor_graph.add_node('extract_receipt', extract_receipt_node)
receipt_processor_graph.add_node('categorize_receipt', categorize_receipt_node)
receipt_processor_graph.add_node('save_receipt', save_receipt_node)
receipt_processor_graph.add_node('compose_confirmation', compose_confirmation_node)
receipt_processor_graph.add_node('handle_error', handle_error_node)

receipt_processor_graph.add_edge(START, 'load_image')
receipt_processor_graph.add_conditional_edges(
    'load_image',
    route_after_image,
    {'extract_receipt': 'extract_receipt', 'handle_error': 'handle_error'},
)
receipt_processor_graph.add_edge('extract_receipt', 'categorize_receipt')
receipt_processor_graph.add_edge('categorize_receipt', 'save_receipt')
receipt_processor_graph.add_conditional_edges(
    'save_receipt',
    route_after_save,
    {'compose_confirmation': 'compose_confirmation', 'handle_error': 'handle_error'},
)
receipt_processor_graph.add_edge('compose_confirmation', END)
receipt_processor_graph.add_edge('handle_error', END)

receipt_processor_app = receipt_processor_graph.compile()

''' Testing '''

if __name__ == '__main__':
    # Process a single receipt photo passed as the first command-line argument.
    if len(sys.argv) > 1:
        result = receipt_processor_app.invoke({'image_path': sys.argv[1]})
        print('Status:', result.get('status'))
        print('Category:', result.get('category'))
        print('Excel file:', result.get('excel_path'))
        for message in result.get('messages', []):
            if isinstance(message, AIMessage):
                print(message.content)
    else:
        print('Usage: python receipt_processor.py <path_to_receipt_photo>')

# --- receipt_processor_prompts.py --- #

EXTRACT_RECEIPT_PROMPT = """
# Role
You are a precise receipt-data extraction specialist that reads photos of purchase receipts.

# Objective
Read the attached receipt photo and convert everything relevant into structured data.

# Inputs
- A photo of a single purchase receipt.
- Today's date is {today} (current year: {year}); use it only to resolve missing years or unreadable dates.

# Instructions
1. Read the printed text of the receipt carefully, including the item table and the totals.
2. location: the store or merchant name; add the branch or city when it is visible on the receipt.
3. date: the receipt date in ISO format (YYYY-MM-DD). If the year is not printed, use {year}. If the date is completely unreadable, use {today}.
4. currency: use the ISO 4217 code when you can recognise the symbol or text ($ -> USD, € -> EUR, £ -> GBP, ₪ -> ILS, ¥ -> JPY); otherwise keep the symbol exactly as printed.
5. total_cost: the final total actually paid (after discounts, including taxes). If no total is printed, sum the item lines.
6. items: include every purchased line item with:
   - name: the item name, cleaned up but faithful to the receipt;
   - quantity: the number of units purchased (default 1 when not shown);
   - unit_price: the price of ONE unit (derive it by dividing the line total by the quantity when only a line total is shown).
7. Exclude non-purchase lines such as payment details, card numbers, loyalty codes, change, and grand totals.

# Hard Rules
- Never invent items that are not visible on the receipt.
- Item names must NOT contain commas or parentheses; replace them with spaces so the records stay machine-parseable.
- All numbers must be plain decimal numbers without currency symbols or thousands separators.
- If the photo is not a receipt or is unreadable, still return the schema with zero/empty values and location "Unknown".

# Output Format
Respond with a ReceiptData object that matches exactly this schema:
- location: string — store or merchant name, with branch or city when visible.
- date: string — ISO "YYYY-MM-DD".
- currency: string — ISO code or the printed symbol.
- total_cost: number — final total paid.
- items: list of objects, each with:
  - name: string
  - quantity: number
  - unit_price: number
"""

CATEGORIZE_RECEIPT_PROMPT = """\\
# Role
You are a personal-finance categorisation assistant.

# Objective
Assign exactly one spending category to a purchase, based primarily on the items bought and secondarily on the location.

# Inputs
- Items: {items}
- Location: {location}

# Instructions
1. Base the decision on what the items are, not on how much they cost.
2. For mixed purchases, choose the category that covers the dominant share of the cost.
3. Choose exactly one category from this list: {categories}.

# Category Guidance
- Groceries: food, drinks, and household supplies from supermarkets, markets, or grocery stores.
- Dining: restaurants, cafés, fast food, takeaway, and delivery.
- Transport: fuel, parking, public transport, taxis, car maintenance and parts.
- Utilities: electricity, water, gas, internet, mobile phone, and other household bills.
- Shopping: clothing, electronics, home goods, cosmetics, gifts, and general retail.
- Health: pharmacies, medicines, medical visits, glasses, and supplements.
- Entertainment: cinema, games, sports events, hobbies, and leisure activities.
- Travel: flights, hotels, trains, car rentals, and vacation-related purchases.
- Education: courses, tuition, school supplies, and textbooks.
- Other: anything that clearly fits none of the above.

# Hard Rules
- Respond with exactly one category, spelled exactly as it appears in the list.
- Do not invent new categories and do not explain your choice.

# Output Format
Respond with a ReceiptCategory object with a single field:
- category: string — one of: {categories}
"""

# --- receipt_storage.py --- #
''' Imports '''

import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

''' Constants '''

load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent.parent / '.env')

# Column headers of the receipts Excel file, in order.
EXCEL_HEADERS = ['cost', 'date', 'items', 'location', 'category']
EXCEL_SHEET_NAME = 'Receipts'
DEFAULT_EXCEL_FILENAME = 'receipts.xlsx'

# Matches entries like "Milk (2 x 3.50 USD)" inside the items column.
ITEM_PATTERN = re.compile(
    r'(?P<name>[^,()]+?)\\s*\\(\\s*(?P<quantity>[\\d.]+)\\s*x\\s*(?P<price>[\\d.]+)\\s*(?P<currency>[^)]*?)\\s*\\)'
)

''' Helpful Functions '''

def get_excel_path() -> Path:
    """Resolve the Excel file path from the environment or the default location."""
    configured = (os.getenv('RECEIPT_EXCEL_PATH') or '').strip()
    if configured:
        return Path(configured).expanduser()
    return Path(__file__).resolve().parent / DEFAULT_EXCEL_FILENAME


def _to_float(value: Any, default: float = 0.0) -> float:
    """Best-effort conversion of spreadsheet values to float."""
    if value is None:
        return default
    try:
        return float(str(value).strip().replace(',', ''))
    except (TypeError, ValueError):
        return default


def _format_number(value: float) -> str:
    """Render numbers compactly: whole values without decimals, else two decimals."""
    if value == int(value):
        return str(int(value))
    return f'{value:.2f}'


def _clean_item_name(name: Any) -> str:
    """Sanitize an item name so the items column stays machine-parseable."""
    text = str(name or '').replace(',', ' ').replace('(', ' ').replace(')', ' ')
    return ' '.join(text.split())


def normalize_date(value: Any) -> str:
    """Normalize a date value to ISO 'YYYY-MM-DD'; returns '' when unparseable."""
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    text = str(value).strip()
    if not text:
        return ''
    try:
        return datetime.fromisoformat(text).strftime('%Y-%m-%d')
    except ValueError:
        pass
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%Y/%m/%d', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(text, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return ''


def format_items_string(receipt: Dict[str, Any]) -> str:
    """Render receipt items as 'item (quantity x price currency), ...'."""
    currency = str(receipt.get('currency') or '').strip()
    parts: List[str] = []
    for item in receipt.get('items') or []:
        if not isinstance(item, dict):
            continue
        name = _clean_item_name(item.get('name', ''))
        if not name:
            continue
        quantity = _to_float(item.get('quantity'), default=1.0)
        unit_price = _to_float(item.get('unit_price'))
        parts.append(f'{name} ({_format_number(quantity)} x {_format_number(unit_price)} {currency})')
    return ', '.join(parts)


def parse_items(items_text: str) -> List[Dict[str, Any]]:
    """Parse the items column back into structured item entries."""
    parsed: List[Dict[str, Any]] = []
    for match in ITEM_PATTERN.finditer(items_text or ''):
        name = ' '.join(match.group('name').split())
        if not name:
            continue
        parsed.append(
            {
                'name': name,
                'quantity': _to_float(match.group('quantity'), default=1.0),
                'unit_price': _to_float(match.group('price')),
                'currency': match.group('currency').strip(),
            }
        )
    return parsed


def append_receipt_record(receipt: Dict[str, Any], category: str) -> Path:
    """Append one categorized receipt as a new row, creating the file if needed.

    The row layout is: cost, date, items, location, category.
    Returns the path of the Excel file that was written.
    """
    path = get_excel_path()
    cost = _to_float(receipt.get('total_cost'))
    date_text = normalize_date(receipt.get('date'))
    items_text = format_items_string(receipt)
    location = str(receipt.get('location') or '').strip() or 'Unknown'
    category_text = str(category or '').strip() or 'Other'

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[EXCEL_SHEET_NAME] if EXCEL_SHEET_NAME in workbook.sheetnames else workbook.active
        if sheet.max_row == 1 and sheet.cell(row=1, column=1).value is None:
            sheet.append(EXCEL_HEADERS)
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = EXCEL_SHEET_NAME
        sheet.append(EXCEL_HEADERS)

    sheet.append([cost, date_text, items_text, location, category_text])
    workbook.save(path)
    workbook.close()
    return path


def load_records() -> List[Dict[str, Any]]:
    """Read every receipt row from the Excel file; empty list when it is missing."""
    path = get_excel_path()
    if not path.exists():
        return []
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[EXCEL_SHEET_NAME] if EXCEL_SHEET_NAME in workbook.sheetnames else workbook.active
        records: List[Dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None or all(value is None for value in row):
                continue
            values = list(row) + [None] * (len(EXCEL_HEADERS) - len(row))
            cost, date_value, items, location, category = values[: len(EXCEL_HEADERS)]
            records.append(
                {
                    'cost': _to_float(cost),
                    'date': normalize_date(date_value),
                    'items': '' if items is None else str(items).strip(),
                    'location': '' if location is None else str(location).strip(),
                    'category': '' if category is None else str(category).strip(),
                }
            )
        return records
    finally:
        workbook.close()


def filter_records_by_date(
    records: List[Dict[str, Any]], start_date: str = '', end_date: str = ''
) -> List[Dict[str, Any]]:
    """Keep records whose ISO date falls inside the optional inclusive range."""
    start = normalize_date(start_date)
    end = normalize_date(end_date)
    if not start and not end:
        return list(records)
    filtered: List[Dict[str, Any]] = []
    for record in records:
        record_date = record.get('date') or ''
        if not record_date:
            continue
        if start and record_date < start:
            continue
        if end and record_date > end:
            continue
        filtered.append(record)
    return filtered