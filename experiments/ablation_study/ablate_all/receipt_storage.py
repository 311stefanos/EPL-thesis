''' Imports '''

import os
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List

from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

''' Constants '''

load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent.parent / '.env')

# Column headers of the receipts Excel file, in order.
EXCEL_HEADERS = ['cost', 'date', 'items', 'location', 'category']
EXCEL_SHEET_NAME = 'Receipts'
DEFAULT_EXCEL_FILENAME = 'receipts.xlsx'

# Matches entries like "Milk (2 x 3.50 USD)" inside the items column.
ITEM_PATTERN = re.compile(
    r'(?P<name>[^,()]+?)\\s*\\(\\s*(?P<quantity>[\\d.]+)\\s*x\\s*(?P<price>[\\d.]+)\\s*(?P<currency>[^)]*?)\\s*\\)'
)

''' Helpful Functions '''

def get_excel_path() -> Path:
    """Resolve the Excel file path from the environment or the default location."""
    configured = (os.getenv('RECEIPT_EXCEL_PATH') or '').strip()
    if configured:
        return Path(configured).expanduser()
    return Path(__file__).resolve().parent / DEFAULT_EXCEL_FILENAME


def _to_float(value: Any, default: float = 0.0) -> float:
    """Best-effort conversion of spreadsheet values to float."""
    if value is None:
        return default
    try:
        return float(str(value).strip().replace(',', ''))
    except (TypeError, ValueError):
        return default


def _format_number(value: float) -> str:
    """Render numbers compactly: whole values without decimals, else two decimals."""
    if value == int(value):
        return str(int(value))
    return f'{value:.2f}'


def _clean_item_name(name: Any) -> str:
    """Sanitize an item name so the items column stays machine-parseable."""
    text = str(name or '').replace(',', ' ').replace('(', ' ').replace(')', ' ')
    return ' '.join(text.split())


def normalize_date(value: Any) -> str:
    """Normalize a date value to ISO 'YYYY-MM-DD'; returns '' when unparseable."""
    if value is None:
        return ''
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    text = str(value).strip()
    if not text:
        return ''
    try:
        return datetime.fromisoformat(text).strftime('%Y-%m-%d')
    except ValueError:
        pass
    for fmt in ('%Y-%m-%d', '%d/%m/%Y', '%m/%d/%Y', '%Y/%m/%d', '%d-%m-%Y', '%d.%m.%Y'):
        try:
            return datetime.strptime(text, fmt).strftime('%Y-%m-%d')
        except ValueError:
            continue
    return ''


def format_items_string(receipt: Dict[str, Any]) -> str:
    """Render receipt items as 'item (quantity x price currency), ...'."""
    currency = str(receipt.get('currency') or '').strip()
    parts: List[str] = []
    for item in receipt.get('items') or []:
        if not isinstance(item, dict):
            continue
        name = _clean_item_name(item.get('name', ''))
        if not name:
            continue
        quantity = _to_float(item.get('quantity'), default=1.0)
        unit_price = _to_float(item.get('unit_price'))
        parts.append(f'{name} ({_format_number(quantity)} x {_format_number(unit_price)} {currency})')
    return ', '.join(parts)


def parse_items(items_text: str) -> List[Dict[str, Any]]:
    """Parse the items column back into structured item entries."""
    parsed: List[Dict[str, Any]] = []
    for match in ITEM_PATTERN.finditer(items_text or ''):
        name = ' '.join(match.group('name').split())
        if not name:
            continue
        parsed.append(
            {
                'name': name,
                'quantity': _to_float(match.group('quantity'), default=1.0),
                'unit_price': _to_float(match.group('price')),
                'currency': match.group('currency').strip(),
            }
        )
    return parsed


def append_receipt_record(receipt: Dict[str, Any], category: str) -> Path:
    """Append one categorized receipt as a new row, creating the file if needed.

    The row layout is: cost, date, items, location, category.
    Returns the path of the Excel file that was written.
    """
    path = get_excel_path()
    cost = _to_float(receipt.get('total_cost'))
    date_text = normalize_date(receipt.get('date'))
    items_text = format_items_string(receipt)
    location = str(receipt.get('location') or '').strip() or 'Unknown'
    category_text = str(category or '').strip() or 'Other'

    if path.exists():
        workbook = load_workbook(path)
        sheet = workbook[EXCEL_SHEET_NAME] if EXCEL_SHEET_NAME in workbook.sheetnames else workbook.active
        if sheet.max_row == 1 and sheet.cell(row=1, column=1).value is None:
            sheet.append(EXCEL_HEADERS)
    else:
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = EXCEL_SHEET_NAME
        sheet.append(EXCEL_HEADERS)

    sheet.append([cost, date_text, items_text, location, category_text])
    workbook.save(path)
    workbook.close()
    return path


def load_records() -> List[Dict[str, Any]]:
    """Read every receipt row from the Excel file; empty list when it is missing."""
    path = get_excel_path()
    if not path.exists():
        return []
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook[EXCEL_SHEET_NAME] if EXCEL_SHEET_NAME in workbook.sheetnames else workbook.active
        records: List[Dict[str, Any]] = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row is None or all(value is None for value in row):
                continue
            values = list(row) + [None] * (len(EXCEL_HEADERS) - len(row))
            cost, date_value, items, location, category = values[: len(EXCEL_HEADERS)]
            records.append(
                {
                    'cost': _to_float(cost),
                    'date': normalize_date(date_value),
                    'items': '' if items is None else str(items).strip(),
                    'location': '' if location is None else str(location).strip(),
                    'category': '' if category is None else str(category).strip(),
                }
            )
        return records
    finally:
        workbook.close()


def filter_records_by_date(
    records: List[Dict[str, Any]], start_date: str = '', end_date: str = ''
) -> List[Dict[str, Any]]:
    """Keep records whose ISO date falls inside the optional inclusive range."""
    start = normalize_date(start_date)
    end = normalize_date(end_date)
    if not start and not end:
        return list(records)
    filtered: List[Dict[str, Any]] = []
    for record in records:
        record_date = record.get('date') or ''
        if not record_date:
            continue
        if start and record_date < start:
            continue
        if end and record_date > end:
            continue
        filtered.append(record)
    return filtered