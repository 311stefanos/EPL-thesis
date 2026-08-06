# LangChain imports
from langchain_core.messages import AIMessage, BaseMessage, HumanMessage, SystemMessage, ToolMessage
from langchain_core.tools import tool

# LangGraph imports
from langgraph.constants import END, START
from langgraph.graph import StateGraph

# Schema imports
from typing import Any, Dict, List, Literal, Optional, Required, Tuple, TypedDict, Union

# General imports
from dotenv import load_dotenv
from pathlib import Path
from io import BytesIO
from urllib.parse import urljoin, urlparse
import base64
import csv
import hashlib
import html
import ipaddress
import json
import mimetypes
import os
import re
import shutil
import socket
import stat
import subprocess
import tempfile
import time
import traceback
import uuid
import zipfile

# Project imports
from utils.utils import myChatOpenAI, safe_invoke, print_function_name, parse_tool_arguments, clean_llm_output
from creations.gaia_solver import gaia_solver_prompts as prompts


''' Constants '''
load_dotenv(dotenv_path=Path(__file__).resolve().parent.parent.parent / ".env")

DEBUG = str(os.getenv("DEBUG", "")).lower() in {"1", "true", "yes", "on"}

BLUE = "\033[94m"
RED = "\033[91m"
GREEN = "\033[92m"
RESET = "\033[0m"

SUPPORTED_ATTACHMENT_EXTENSIONS = {
    "csv", "xlsx", "json", "xml", "docx", "pptx", "pdf", "txt", "py", "zip",
    "png", "jpg", "jpeg", "webp", "gif", "bmp", "tiff",
}

TOOL_CATEGORY_TO_NAME = {
    "web_search": "web_search_tool",
    "webpage_read": "webpage_read_tool",
    "deterministic_parser": "deterministic_parser_tool",
    "restricted_python": "restricted_python_tool",
    "multimodal_inspection": "multimodal_inspection_tool",
    "secure_zip": "secure_zip_tool",
}
TOOL_NAME_TO_CATEGORY = {
    tool_name: category
    for category, tool_name in TOOL_CATEGORY_TO_NAME.items()
}

MANAGED_TEMP_DIR = Path(
    os.getenv(
        "GAIA_MANAGED_TEMP_DIR",
        str(Path(tempfile.gettempdir()) / "gaia_solver"),
    )
).resolve()
MANAGED_TEMP_DIR.mkdir(parents=True, exist_ok=True)

MAX_TOOL_OUTPUT_CHARS = int(os.getenv("GAIA_MAX_TOOL_OUTPUT_CHARS", "50000"))
MAX_WEB_RESPONSE_BYTES = int(os.getenv("GAIA_MAX_WEB_RESPONSE_BYTES", str(15 * 1024 * 1024)))
MAX_LOCAL_FILE_BYTES = int(os.getenv("GAIA_MAX_LOCAL_FILE_BYTES", str(50 * 1024 * 1024)))
MAX_ZIP_ENTRIES = int(os.getenv("GAIA_MAX_ZIP_ENTRIES", "500"))
MAX_ZIP_MEMBER_BYTES = int(os.getenv("GAIA_MAX_ZIP_MEMBER_BYTES", str(50 * 1024 * 1024)))
MAX_ZIP_TOTAL_BYTES = int(os.getenv("GAIA_MAX_ZIP_TOTAL_BYTES", str(200 * 1024 * 1024)))
MAX_ZIP_RATIO = float(os.getenv("GAIA_MAX_ZIP_RATIO", "100"))
MAX_WEB_RESULTS = int(os.getenv("GAIA_MAX_WEB_RESULTS", "10"))

_PERMITTED_FILES: set[Path] = set()
_PERMITTED_ROOTS: set[Path] = {MANAGED_TEMP_DIR, Path.cwd().resolve()}

print(f"\n{BLUE}[AGENT] [INFO] [STARTUP]{RESET} gaia_solver") if DEBUG else None


# Schemas
class AnswerRequirements(TypedDict):
    answer_type: str
    requested_output_format: str
    separator: Optional[str]
    expected_item_count: Optional[int]
    ordering: Optional[str]
    units: Optional[str]
    rounding: Optional[Union[int, str]]
    capitalization: Optional[str]
    date_format: Optional[str]
    date_restrictions: Optional[str]
    required_sources: List[str]
    prohibited_sources: List[str]
    filtering_rules: List[str]
    required_calculations: List[str]
    required_prefix: Optional[str]
    required_suffix: Optional[str]
    prohibited_characters: List[str]
    prohibited_extra_text: List[str]


class Observation(TypedDict):
    observation_id: str
    content: str
    value: Any
    source: str
    source_location: str
    tool_name: str
    confidence: float
    uncertainty: Union[str, float]


class Calculation(TypedDict):
    calculation_id: str
    description: str
    inputs: Dict[str, Any]
    method: str
    result: Any
    units: str
    completed: bool
    error: Optional[str]


class ToolCallLogEntry(TypedDict):
    tool_name: str
    arguments: Dict[str, Any]
    start_time: float
    end_time: float
    status: str
    error: Optional[str]


class PlanStep(TypedDict):
    step_id: Union[str, int]
    objective: str
    tool_category: Literal[
        "web_search",
        "webpage_read",
        "deterministic_parser",
        "restricted_python",
        "multimodal_inspection",
        "secure_zip",
    ]
    focused_instruction: str
    required_inputs: Dict[str, Any]
    expected_result: str
    success_criteria: str
    fallback_action: Literal["retry", "skip_optional", "replan", "fail"]
    required: bool
    status: Literal["pending", "completed", "failed", "skipped"]
    attempts: int


class AnswerSupportEntry(TypedDict):
    answer_part: str
    supporting_observations: List[str]
    supporting_calculations: List[str]


class AgentSchema(TypedDict, total=False):
    question: Required[str]
    messages: List[BaseMessage]
    attachment_path: Optional[str]
    attachment_extension: Optional[str]
    attachment_supported: Optional[bool]
    answer_requirements: AnswerRequirements
    observations: List[Observation]
    calculations: List[Calculation]
    tool_call_log: List[ToolCallLogEntry]
    errors: List[Dict[str, Any]]
    execution_limits: Dict[str, Any]
    execution_counters: Dict[str, Any]
    reviewer_feedback: Optional[str]
    failed_steps: List[Dict[str, Any]]
    plan_history: List[List[PlanStep]]
    current_step_index: int
    current_plan: List[PlanStep]
    status: Literal["running", "failed", "success"]
    failure_reason: Optional[str]
    next_action: Literal["execute_plan", "create_plan", "synthesize_candidate", "__end__"]
    candidate_answer: str
    answer_support: List[AnswerSupportEntry]
    unresolved_requirements: List[str]
    unsupported_components: List[str]
    reviewer_decision: Literal[
        "approve",
        "revise_answer",
        "recalculate",
        "gather_more_evidence",
        "revise_plan",
    ]
    format_check_summary: str
    final_output: Optional[str]
    started_at: float


# Helpful Functions
def _ensure_state_defaults(state: AgentSchema) -> None:
    state.setdefault("messages", [])
    state.setdefault("observations", [])
    state.setdefault("calculations", [])
    state.setdefault("tool_call_log", [])
    state.setdefault("errors", [])
    state.setdefault("execution_limits", {
        "total_tool_calls": 30,
        "max_step_attempts": 3,
        "max_repeated_searches": 4,
        "max_duplicate_actions": 4,
        "max_replans": 3,
        "total_task_runtime": 600,
    })
    state.setdefault("execution_counters", {
        "total_tool_calls": 0,
        "repeated_searches": 0,
        "duplicate_actions": 0,
        "replans": 0,
    })
    state.setdefault("failed_steps", [])
    state.setdefault("plan_history", [])
    state.setdefault("current_plan", [])
    state.setdefault("current_step_index", 0)
    state.setdefault("status", "running")
    state.setdefault("failure_reason", None)
    state.setdefault("reviewer_feedback", None)
    state.setdefault("started_at", time.time())


def _json_safe(value: Any) -> Any:
    if isinstance(value, BaseMessage):
        return {
            "type": value.__class__.__name__,
            "content": value.content,
            "tool_calls": getattr(value, "tool_calls", None),
        }
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, bytes):
        return f"<{len(value)} bytes>"
    try:
        json.dumps(value)
        return value
    except TypeError:
        return str(value)


def _as_prompt_text(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=_json_safe, indent=2)


def _extract_json_response(result: Any) -> Any:
    if isinstance(result, (dict, list)):
        return result
    if hasattr(result, "model_dump"):
        return result.model_dump()

    content = result.content if hasattr(result, "content") else str(result)
    content = clean_llm_output(str(content)).strip()

    try:
        return json.loads(content)
    except json.JSONDecodeError:
        decoder = json.JSONDecoder()
        for index, character in enumerate(content):
            if character not in "[{":
                continue
            try:
                parsed, _ = decoder.raw_decode(content[index:])
                return parsed
            except json.JSONDecodeError:
                continue
    raise ValueError("The LLM response did not contain valid JSON.")


def _get_tool_calls(message: BaseMessage) -> List[Dict[str, Any]]:
    calls = getattr(message, "tool_calls", None)
    if calls:
        return list(calls)

    raw_calls = getattr(message, "additional_kwargs", {}).get("tool_calls", [])
    parsed_calls: List[Dict[str, Any]] = []
    for raw_call in raw_calls:
        function = raw_call.get("function", {})
        parsed_calls.append({
            "id": raw_call.get("id"),
            "name": function.get("name"),
            "args": parse_tool_arguments(function.get("arguments", "{}")),
        })
    return parsed_calls


def _register_attachment_path(attachment_path: Optional[str]) -> None:
    if not attachment_path:
        return
    try:
        path = Path(attachment_path).expanduser().resolve()
        _PERMITTED_FILES.add(path)
    except OSError:
        return


def _is_relative_to(path: Path, root: Path) -> bool:
    try:
        path.relative_to(root)
        return True
    except ValueError:
        return False


def _validate_local_file(file_path: str) -> Path:
    path = Path(file_path).expanduser().resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"File does not exist: {path}")
    if path.stat().st_size > MAX_LOCAL_FILE_BYTES:
        raise ValueError(f"File exceeds the configured size limit: {path}")
    if path not in _PERMITTED_FILES and not any(_is_relative_to(path, root) for root in _PERMITTED_ROOTS):
        raise PermissionError(f"File is outside the permitted directories: {path}")
    return path


def _safe_temp_path(suffix: str = "", prefix: str = "gaia_") -> Path:
    safe_suffix = suffix if suffix.startswith(".") or suffix == "" else f".{suffix}"
    path = MANAGED_TEMP_DIR / f"{prefix}{uuid.uuid4().hex}{safe_suffix}"
    return path.resolve()


def _validate_public_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"}:
        raise ValueError("Only HTTP and HTTPS URLs are allowed.")
    if not parsed.hostname:
        raise ValueError("URL does not contain a hostname.")

    hostname = parsed.hostname.lower()
    if hostname == "localhost" or hostname.endswith(".localhost"):
        raise ValueError("Localhost URLs are not allowed.")

    try:
        addresses = socket.getaddrinfo(hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    except socket.gaierror as exc:
        raise ValueError(f"Could not resolve URL hostname: {hostname}") from exc

    for address in addresses:
        ip = ipaddress.ip_address(address[4][0])
        if (
            ip.is_private
            or ip.is_loopback
            or ip.is_link_local
            or ip.is_multicast
            or ip.is_reserved
            or ip.is_unspecified
        ):
            raise ValueError(f"URL resolves to a prohibited address: {ip}")


def _bounded_text(text: str, limit: int = MAX_TOOL_OUTPUT_CHARS) -> Tuple[str, bool]:
    if len(text) <= limit:
        return text, False
    return text[:limit], True


def _select_passages(text: str, instruction: str, maximum: int = 8) -> List[Dict[str, str]]:
    paragraphs = [
        re.sub(r"\s+", " ", paragraph).strip()
        for paragraph in re.split(r"\n\s*\n", text)
        if paragraph.strip()
    ]
    if not paragraphs:
        return []

    terms = {
        term.lower()
        for term in re.findall(r"[A-Za-z0-9À-ž_-]{3,}", instruction)
    }
    scored = []
    offset = 0
    for paragraph in paragraphs:
        lowered = paragraph.lower()
        score = sum(lowered.count(term) for term in terms)
        scored.append((score, offset, paragraph))
        offset += len(paragraph) + 2

    if terms and any(score > 0 for score, _, _ in scored):
        selected = sorted(scored, key=lambda item: (-item[0], item[1]))[:maximum]
        selected.sort(key=lambda item: item[1])
    else:
        selected = scored[:maximum]

    return [
        {
            "content": paragraph[:5000],
            "source_location": f"character offset {offset}",
        }
        for _, offset, paragraph in selected
    ]


def _normalise_tool_call(call: Dict[str, Any]) -> Tuple[str, Dict[str, Any], str]:
    name = call.get("name")
    args = call.get("args", {})
    call_id = call.get("id") or f"call_{uuid.uuid4().hex}"

    if not isinstance(name, str) or not name:
        raise ValueError("Tool call is missing a valid tool name.")
    if isinstance(args, str):
        args = parse_tool_arguments(args)
    if not isinstance(args, dict):
        raise ValueError("Tool arguments must be a dictionary.")
    return name, args, call_id


def _limit_failure(state: AgentSchema) -> Optional[str]:
    limits = state["execution_limits"]
    counters = state["execution_counters"]

    checks = (
        ("total_tool_calls", "total_tool_calls", "Total tool-call limit reached"),
        ("max_repeated_searches", "repeated_searches", "Repeated-search limit reached"),
        ("max_duplicate_actions", "duplicate_actions", "Duplicate-action limit reached"),
    )
    for limit_key, counter_key, message in checks:
        limit = limits.get(limit_key)
        if limit is not None and counters.get(counter_key, 0) >= limit:
            return message

    runtime_limit = limits.get("total_task_runtime")
    if runtime_limit is not None and time.time() - state["started_at"] >= runtime_limit:
        return "Total task runtime limit reached"

    return None


def _route_failed_step(state: AgentSchema, step: PlanStep, reason: str) -> None:
    step["status"] = "failed"
    state["failed_steps"].append({
        "step_id": step.get("step_id"),
        "reason": reason,
        "attempts": step.get("attempts", 0),
    })

    fallback = step.get("fallback_action", "fail")
    attempts = step.get("attempts", 0)
    max_attempts = state["execution_limits"].get("max_step_attempts", 3)

    if fallback == "retry" and attempts < max_attempts:
        state["next_action"] = "execute_plan"
        return

    if fallback == "skip_optional" and not step.get("required", True):
        step["status"] = "skipped"
        next_index = state["current_step_index"] + 1
        state["current_step_index"] = next_index
        state["next_action"] = (
            "execute_plan"
            if next_index < len(state["current_plan"])
            else "synthesize_candidate"
        )
        return

    if fallback == "replan":
        max_replans = state["execution_limits"].get("max_replans", 3)
        if state["execution_counters"].get("replans", 0) < max_replans:
            state["next_action"] = "create_plan"
            return

    state["status"] = "failed"
    state["failure_reason"] = reason
    state["next_action"] = "__end__"


def get_attachment_info(attachment_path: Optional[str]) -> Tuple[Optional[str], Optional[bool]]:
    if not attachment_path:
        return None, None
    extension = Path(attachment_path).suffix.lower().lstrip(".")
    return extension, extension in SUPPORTED_ATTACHMENT_EXTENSIONS


# Tools
@tool
def web_search_tool(query: str, max_results: int = 5) -> List[Dict[str, str]]:
    """Search the web with Tavily and return title, URL, and snippet fields."""
    query = query.strip()
    if not query:
        raise ValueError("Search query cannot be empty.")
    if len(query) > 1000:
        raise ValueError("Search query is too long.")

    max_results = max(1, min(int(max_results), MAX_WEB_RESULTS))
    api_key = os.getenv("TAVILY_API_KEY")
    if not api_key:
        raise RuntimeError("TAVILY_API_KEY is not configured.")

    try:
        from tavily import TavilyClient
    except ImportError as exc:
        raise RuntimeError("The 'tavily-python' package is required.") from exc

    response = TavilyClient(api_key=api_key).search(
        query=query,
        max_results=max_results,
        search_depth="advanced",
        include_answer=False,
        include_raw_content=False,
    )

    results = response.get("results", []) if isinstance(response, dict) else []
    return [
        {
            "title": str(item.get("title", "")),
            "url": str(item.get("url", "")),
            "snippet": str(item.get("content", item.get("snippet", ""))),
        }
        for item in results[:max_results]
        if isinstance(item, dict) and item.get("url")
    ]


@tool
def webpage_read_tool(url: str, instruction: str) -> Dict[str, Any]:
    """Read bounded public web content or download a supported public file."""
    try:
        import requests
    except ImportError as exc:
        raise RuntimeError("The 'requests' package is required.") from exc

    _validate_public_url(url)
    session = requests.Session()
    current_url = url
    response = None

    for _ in range(6):
        _validate_public_url(current_url)
        response = session.get(
            current_url,
            timeout=(10, 30),
            allow_redirects=False,
            stream=True,
            headers={"User-Agent": "GAIA-Solver/1.0"},
        )
        if response.is_redirect or response.is_permanent_redirect:
            location = response.headers.get("Location")
            if not location:
                raise RuntimeError("Redirect response did not include a location.")
            current_url = urljoin(current_url, location)
            continue
        break
    else:
        raise RuntimeError("Maximum redirect count exceeded.")

    if response is None:
        raise RuntimeError("No HTTP response was received.")
    response.raise_for_status()

    chunks: List[bytes] = []
    total = 0
    for chunk in response.iter_content(chunk_size=65536):
        if not chunk:
            continue
        total += len(chunk)
        if total > MAX_WEB_RESPONSE_BYTES:
            raise ValueError("Web response exceeds the configured size limit.")
        chunks.append(chunk)
    content = b"".join(chunks)

    content_type = response.headers.get("Content-Type", "").split(";")[0].strip().lower()
    final_url = response.url or current_url
    title = Path(urlparse(final_url).path).name or final_url

    downloadable_mimes = {
        "text/csv", "application/vnd.ms-excel",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/json", "application/xml", "text/xml",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "application/zip",
    }
    is_media = content_type.startswith(("image/", "audio/", "video/"))

    if content_type in downloadable_mimes or is_media:
        guessed_extension = mimetypes.guess_extension(content_type) or Path(urlparse(final_url).path).suffix
        destination = _safe_temp_path(suffix=guessed_extension or ".bin", prefix="download_")
        destination.write_bytes(content)
        _PERMITTED_FILES.add(destination)
        return {
            "title": title,
            "url": final_url,
            "content_type": content_type,
            "passages": [],
            "downloaded_path": str(destination),
            "error": None,
        }

    if content_type == "application/pdf" or final_url.lower().endswith(".pdf"):
        try:
            from pypdf import PdfReader
            reader = PdfReader(BytesIO(content))
            page_text = []
            for index, page in enumerate(reader.pages):
                extracted = page.extract_text() or ""
                page_text.append(f"[Page {index + 1}]\n{extracted}")
            text = "\n\n".join(page_text).strip()
        except Exception:
            destination = _safe_temp_path(suffix=".pdf", prefix="pdf_")
            destination.write_bytes(content)
            _PERMITTED_FILES.add(destination)
            return {
                "title": title,
                "url": final_url,
                "content_type": "application/pdf",
                "passages": [],
                "downloaded_path": str(destination),
                "error": None,
            }
    else:
        encoding = response.encoding or "utf-8"
        decoded = content.decode(encoding, errors="replace")
        if content_type == "text/html" or "<html" in decoded[:1000].lower():
            try:
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(decoded, "html.parser")
                for element in soup(["script", "style", "nav", "header", "footer", "aside"]):
                    element.decompose()
                page_title = soup.title.string.strip() if soup.title and soup.title.string else ""
                title = page_title or title
                text = soup.get_text("\n")
            except ImportError:
                text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", decoded)
                text = re.sub(r"(?s)<[^>]+>", "\n", text)
                text = html.unescape(text)
        else:
            text = decoded

    text, _ = _bounded_text(text, MAX_TOOL_OUTPUT_CHARS * 4)
    return {
        "title": title,
        "url": final_url,
        "content_type": content_type or "text/plain",
        "passages": _select_passages(text, instruction),
        "downloaded_path": None,
        "error": None,
    }


@tool
def deterministic_parser_tool(file_path: str, instruction: str) -> Dict[str, Any]:
    """Parse a supported local file deterministically and return bounded items."""
    path = _validate_local_file(file_path)
    extension = path.suffix.lower().lstrip(".")
    metadata: Dict[str, Any] = {}
    items: List[Dict[str, Any]] = []
    truncated = False
    maximum_items = 250

    try:
        if extension == "csv":
            with path.open("r", encoding="utf-8-sig", newline="") as file:
                reader = csv.reader(file)
                rows = list(reader)
            metadata = {
                "rows": len(rows),
                "columns": max((len(row) for row in rows), default=0),
            }
            for row_index, row in enumerate(rows[:maximum_items], start=1):
                items.append({
                    "content": row,
                    "source_location": f"row {row_index}",
                })
            truncated = len(rows) > maximum_items

        elif extension == "xlsx":
            try:
                from openpyxl import load_workbook
            except ImportError as exc:
                raise RuntimeError("The 'openpyxl' package is required for XLSX files.") from exc
            workbook = load_workbook(path, read_only=True, data_only=True)
            metadata["sheets"] = workbook.sheetnames
            for sheet in workbook.worksheets:
                for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                    items.append({
                        "content": list(row),
                        "source_location": f"sheet {sheet.title}, row {row_index}",
                    })
                    if len(items) >= maximum_items:
                        truncated = True
                        break
                if truncated:
                    break

        elif extension == "json":
            data = json.loads(path.read_text(encoding="utf-8"))
            metadata["top_level_type"] = type(data).__name__
            if isinstance(data, dict):
                metadata["top_level_keys"] = list(data.keys())[:100]
                iterable = list(data.items())
                for key, value in iterable[:maximum_items]:
                    items.append({
                        "content": value,
                        "source_location": f"$.{key}",
                    })
                truncated = len(iterable) > maximum_items
            elif isinstance(data, list):
                for index, value in enumerate(data[:maximum_items]):
                    items.append({
                        "content": value,
                        "source_location": f"$[{index}]",
                    })
                truncated = len(data) > maximum_items
            else:
                items.append({"content": data, "source_location": "$"})

        elif extension == "xml":
            import xml.etree.ElementTree as ET
            root = ET.parse(path).getroot()
            metadata["root_tag"] = root.tag
            for index, element in enumerate(root.iter()):
                items.append({
                    "content": {
                        "tag": element.tag,
                        "attributes": element.attrib,
                        "text": (element.text or "").strip(),
                    },
                    "source_location": f"element {index}: {element.tag}",
                })
                if len(items) >= maximum_items:
                    truncated = True
                    break

        elif extension == "docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("The 'python-docx' package is required for DOCX files.") from exc
            document = Document(path)
            metadata["paragraphs"] = len(document.paragraphs)
            for index, paragraph in enumerate(document.paragraphs[:maximum_items], start=1):
                if paragraph.text.strip():
                    items.append({
                        "content": paragraph.text,
                        "source_location": f"paragraph {index}",
                    })
            truncated = len(document.paragraphs) > maximum_items

        elif extension == "pptx":
            try:
                from pptx import Presentation
            except ImportError as exc:
                raise RuntimeError("The 'python-pptx' package is required for PPTX files.") from exc
            presentation = Presentation(path)
            metadata["slides"] = len(presentation.slides)
            for slide_index, slide in enumerate(presentation.slides, start=1):
                text_parts = [
                    shape.text
                    for shape in slide.shapes
                    if hasattr(shape, "text") and shape.text.strip()
                ]
                items.append({
                    "content": "\n".join(text_parts),
                    "source_location": f"slide {slide_index}",
                })
                if len(items) >= maximum_items:
                    truncated = True
                    break

        elif extension == "pdf":
            try:
                from pypdf import PdfReader
            except ImportError as exc:
                raise RuntimeError("The 'pypdf' package is required for PDF files.") from exc
            reader = PdfReader(path)
            metadata["pages"] = len(reader.pages)
            extracted_any = False
            for page_index, page in enumerate(reader.pages, start=1):
                text = page.extract_text() or ""
                if text.strip():
                    extracted_any = True
                    items.append({
                        "content": text,
                        "source_location": f"page {page_index}",
                    })
                if len(items) >= maximum_items:
                    truncated = True
                    break
            if not extracted_any:
                return {
                    "file_path": str(path),
                    "file_type": extension,
                    "metadata": metadata,
                    "items": [],
                    "truncated": False,
                    "error": "PDF contains no machine-readable text; use multimodal_inspection_tool.",
                }

        elif extension in {"txt", "py"}:
            text = path.read_text(encoding="utf-8", errors="replace")
            lines = text.splitlines()
            metadata["lines"] = len(lines)
            chunk_size = 40
            for start in range(0, min(len(lines), maximum_items * chunk_size), chunk_size):
                chunk = "\n".join(lines[start:start + chunk_size])
                items.append({
                    "content": chunk,
                    "source_location": f"lines {start + 1}-{min(start + chunk_size, len(lines))}",
                })
            truncated = len(lines) > maximum_items * chunk_size

        else:
            return {
                "file_path": str(path),
                "file_type": extension,
                "metadata": {},
                "items": [],
                "truncated": False,
                "error": f"Unsupported deterministic file type: {extension}",
            }

        instruction_terms = {
            term.lower()
            for term in re.findall(r"[A-Za-z0-9À-ž_-]{3,}", instruction)
        }
        if instruction_terms:
            filtered = [
                item
                for item in items
                if any(
                    term in json.dumps(item["content"], ensure_ascii=False, default=str).lower()
                    for term in instruction_terms
                )
            ]
            if filtered:
                items = filtered[:maximum_items]

        return {
            "file_path": str(path),
            "file_type": extension,
            "metadata": metadata,
            "items": items,
            "truncated": truncated,
            "error": None,
        }
    except Exception as exc:
        return {
            "file_path": str(path),
            "file_type": extension,
            "metadata": metadata,
            "items": [],
            "truncated": False,
            "error": str(exc),
        }


@tool
def restricted_python_tool(code: str) -> Dict[str, Any]:
    """Execute Python in a restricted disposable Docker container."""
    if not code.strip():
        raise ValueError("Python code cannot be empty.")
    if len(code) > 100000:
        raise ValueError("Python code exceeds the configured length limit.")
    if shutil.which("docker") is None:
        raise RuntimeError("Docker is not installed or is not available on PATH.")

    run_directory = Path(tempfile.mkdtemp(prefix="python_", dir=MANAGED_TEMP_DIR))
    script_path = run_directory / "main.py"
    script_path.write_text(code, encoding="utf-8")

    timeout_seconds = int(os.getenv("GAIA_PYTHON_TIMEOUT_SECONDS", "30"))
    command = [
        "docker", "run", "--rm",
        "--network", "none",
        "--read-only",
        "--user", "65534:65534",
        "--cap-drop", "ALL",
        "--security-opt", "no-new-privileges",
        "--memory", os.getenv("GAIA_DOCKER_MEMORY", "256m"),
        "--cpus", os.getenv("GAIA_DOCKER_CPUS", "1"),
        "--pids-limit", os.getenv("GAIA_DOCKER_PIDS", "64"),
        "--tmpfs", "/tmp:rw,noexec,nosuid,size=64m",
        "-v", f"{script_path.resolve()}:/workspace/main.py:ro",
        os.getenv("GAIA_DOCKER_IMAGE", "python:3.11-slim"),
        "python", "/workspace/main.py",
    ]

    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=timeout_seconds,
            check=False,
        )
        stdout, _ = _bounded_text(completed.stdout)
        stderr, _ = _bounded_text(completed.stderr)
        succeeded = completed.returncode == 0
        return {
            "stdout": stdout,
            "stderr": stderr,
            "exit_code": completed.returncode,
            "timed_out": False,
            "completed": succeeded,
            "error": None if succeeded else (stderr.strip() or f"Process exited with code {completed.returncode}"),
        }
    except subprocess.TimeoutExpired as exc:
        stdout = exc.stdout.decode() if isinstance(exc.stdout, bytes) else (exc.stdout or "")
        stderr = exc.stderr.decode() if isinstance(exc.stderr, bytes) else (exc.stderr or "")
        stdout, _ = _bounded_text(stdout)
        stderr, _ = _bounded_text(stderr)
        return {
            "stdout": stdout,
            "stderr": stderr,
            "exit_code": None,
            "timed_out": True,
            "completed": False,
            "error": f"Execution exceeded {timeout_seconds} seconds.",
        }
    finally:
        shutil.rmtree(run_directory, ignore_errors=True)


@tool
def multimodal_inspection_tool(file_path: str, instruction: str) -> List[Dict[str, Any]]:
    """Inspect an image or image-based PDF with the configured multimodal model."""
    path = _validate_local_file(file_path)
    mime_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    content_parts: List[Dict[str, Any]] = [{
        "type": "text",
        "text": (
            f"{instruction}\n\n"
            "Return only a JSON array. Each item must contain content, value, "
            "source_location, confidence, and uncertainty. Extract evidence only."
        ),
    }]

    if mime_type.startswith("image/"):
        encoded = base64.b64encode(path.read_bytes()).decode("ascii")
        content_parts.append({
            "type": "image_url",
            "image_url": {"url": f"data:{mime_type};base64,{encoded}"},
        })
    elif path.suffix.lower() == ".pdf":
        try:
            import fitz
        except ImportError as exc:
            raise RuntimeError("The 'PyMuPDF' package is required for scanned PDFs.") from exc
        document = fitz.open(path)
        max_pages = min(len(document), int(os.getenv("GAIA_MULTIMODAL_MAX_PAGES", "8")))
        for page_index in range(max_pages):
            pixmap = document[page_index].get_pixmap(matrix=fitz.Matrix(1.5, 1.5), alpha=False)
            encoded = base64.b64encode(pixmap.tobytes("png")).decode("ascii")
            content_parts.append({
                "type": "text",
                "text": f"PDF page {page_index + 1}",
            })
            content_parts.append({
                "type": "image_url",
                "image_url": {"url": f"data:image/png;base64,{encoded}"},
            })
    else:
        raise ValueError(
            "This implementation supports images and scanned PDFs. "
            "Audio and video require a provider-specific adapter."
        )

    provider = os.getenv("MULTIMODAL_PROVIDER") or os.getenv("PROVIDER")
    model = os.getenv("MULTIMODAL_MODEL", "google/gemini-2.5-flash-lite")
    multimodal_llm = myChatOpenAI(provider=provider, model=model, temperature=0.0)
    result = safe_invoke(
        multimodal_llm,
        messages=[HumanMessage(content=content_parts)],
    )
    parsed = _extract_json_response(result)
    if not isinstance(parsed, list):
        raise ValueError("Multimodal model did not return a JSON list.")

    observations = []
    for item in parsed:
        if not isinstance(item, dict):
            continue
        observations.append({
            "content": str(item.get("content", "")),
            "value": item.get("value"),
            "source_location": str(item.get("source_location", "unknown")),
            "confidence": float(item.get("confidence", 0.5) or 0.5),
            "uncertainty": str(item.get("uncertainty", "unspecified")),
        })
    return observations


@tool
def secure_zip_tool(
    zip_path: str,
    action: Literal["list", "extract"],
    member_path: Optional[str] = None,
) -> Dict[str, Any]:
    """Validate a ZIP archive, list members, or extract one regular file."""
    path = _validate_local_file(zip_path)
    if action == "list" and member_path is not None:
        return {
            "action": action,
            "members": [],
            "extracted_path": None,
            "truncated": False,
            "error": "member_path must be omitted when action is 'list'.",
        }
    if action == "extract" and not member_path:
        return {
            "action": action,
            "members": [],
            "extracted_path": None,
            "truncated": False,
            "error": "member_path is required when action is 'extract'.",
        }

    with zipfile.ZipFile(path) as archive:
        infos = archive.infolist()
        if len(infos) > MAX_ZIP_ENTRIES:
            raise ValueError("ZIP archive contains too many entries.")

        total_size = 0
        safe_infos: List[zipfile.ZipInfo] = []
        names_seen: set[str] = set()

        for info in infos:
            name = info.filename.replace("\\", "/")
            pure = Path(name)

            if (
                pure.is_absolute()
                or ".." in pure.parts
                or name.startswith("/")
                or re.match(r"^[A-Za-z]:", name)
            ):
                raise ValueError(f"Unsafe ZIP member path: {name}")

            mode = info.external_attr >> 16
            if stat.S_ISLNK(mode) or stat.S_ISCHR(mode) or stat.S_ISBLK(mode) or stat.S_ISFIFO(mode):
                raise ValueError(f"Unsafe ZIP member type: {name}")
            if info.flag_bits & 0x1:
                raise ValueError("Encrypted ZIP archives are not supported.")
            if info.file_size > MAX_ZIP_MEMBER_BYTES:
                raise ValueError(f"ZIP member exceeds the size limit: {name}")

            compressed = max(info.compress_size, 1)
            if info.file_size / compressed > MAX_ZIP_RATIO:
                raise ValueError(f"ZIP member exceeds the compression-ratio limit: {name}")

            total_size += info.file_size
            if total_size > MAX_ZIP_TOTAL_BYTES:
                raise ValueError("ZIP archive exceeds the total uncompressed-size limit.")

            if name in names_seen:
                raise ValueError(f"ZIP archive contains duplicate member name: {name}")
            names_seen.add(name)
            safe_infos.append(info)

        member_metadata = [
            {
                "name": info.filename,
                "size": info.file_size,
                "compressed_size": info.compress_size,
                "is_directory": info.is_dir(),
            }
            for info in safe_infos[:200]
        ]
        truncated = len(safe_infos) > 200

        if action == "list":
            return {
                "action": action,
                "members": member_metadata,
                "extracted_path": None,
                "truncated": truncated,
                "error": None,
            }

        matching = [info for info in safe_infos if info.filename == member_path]
        if len(matching) != 1:
            return {
                "action": action,
                "members": [],
                "extracted_path": None,
                "truncated": False,
                "error": "Requested member was not found exactly once.",
            }

        info = matching[0]
        if info.is_dir():
            return {
                "action": action,
                "members": [],
                "extracted_path": None,
                "truncated": False,
                "error": "Directories cannot be extracted.",
            }

        suffix = Path(info.filename).suffix
        destination = _safe_temp_path(suffix=suffix, prefix="zip_")
        with archive.open(info, "r") as source, destination.open("wb") as output:
            shutil.copyfileobj(source, output, length=65536)

        _PERMITTED_FILES.add(destination)
        return {
            "action": action,
            "members": [{
                "name": info.filename,
                "size": info.file_size,
                "compressed_size": info.compress_size,
                "is_directory": False,
            }],
            "extracted_path": str(destination),
            "truncated": False,
            "error": None,
        }


TOOLS = [
    web_search_tool,
    webpage_read_tool,
    deterministic_parser_tool,
    restricted_python_tool,
    multimodal_inspection_tool,
    secure_zip_tool,
]
TOOLS_BY_NAME = {tool_object.name: tool_object for tool_object in TOOLS}
TOOLS_BY_CATEGORY = {
    category: TOOLS_BY_NAME[name]
    for category, name in TOOL_CATEGORY_TO_NAME.items()
}


# LLM
analyze_task_llm = myChatOpenAI(
    temperature=0.0,
).with_structured_output(AnswerRequirements)

create_plan_llm = myChatOpenAI(temperature=0.0)
execute_plan_llm = myChatOpenAI(temperature=0.0)
synthesize_candidate_llm = myChatOpenAI(temperature=0.0)
review_candidate_llm = myChatOpenAI(temperature=0.0)
format_output_llm = myChatOpenAI(temperature=0.0)


# Nodes
def analyze_task(state: AgentSchema) -> AgentSchema:
    """Analyze the question and extract output requirements."""
    print_function_name()
    _ensure_state_defaults(state)
    try:
        attachment_path = state.get("attachment_path")
        _register_attachment_path(attachment_path)
        extension, supported = get_attachment_info(attachment_path)

        prompt = prompts.ANALYZE_TASK_PROMPT.format(
            question=state["question"],
            attachment_path=attachment_path or "None",
        )
        result = safe_invoke(
            analyze_task_llm,
            messages=[SystemMessage(content=prompt)],
        )
        requirements = _extract_json_response(result)
        if not isinstance(requirements, dict):
            raise ValueError("analyze_task did not return an AnswerRequirements dictionary.")

        state["answer_requirements"] = requirements
        state["attachment_extension"] = extension
        state["attachment_supported"] = supported
        state["status"] = "running"
        state["failure_reason"] = None
        return state
    except Exception as exc:
        print(f"{RED}[NODE] [ERR]{RESET}", exc) if DEBUG else None
        traceback.print_exc() if DEBUG else None
        state["status"] = "failed"
        state["failure_reason"] = f"analyze_task failed: {exc}"
        return state


def create_plan(state: AgentSchema) -> AgentSchema:
    """Create or revise a bounded structured execution plan."""
    print_function_name()
    _ensure_state_defaults(state)
    if state.get("status") == "failed":
        state["next_action"] = "__end__"
        return state

    try:
        plan_history = state["plan_history"]
        execution_counters = state["execution_counters"]

        prompt = prompts.CREATE_PLAN_PROMPT.format(
            question=state.get("question", ""),
            answer_requirements=_as_prompt_text(state.get("answer_requirements", {})),
            attachment_path=state.get("attachment_path"),
            attachment_extension=state.get("attachment_extension"),
            attachment_supported=state.get("attachment_supported"),
            observations=_as_prompt_text(state.get("observations", [])),
            calculations=_as_prompt_text(state.get("calculations", [])),
            tool_call_log=_as_prompt_text(state.get("tool_call_log", [])),
            errors=_as_prompt_text(state.get("errors", [])),
            execution_limits=_as_prompt_text(state.get("execution_limits", {})),
            execution_counters=_as_prompt_text(execution_counters),
            reviewer_feedback=state.get("reviewer_feedback") or "",
            failed_steps=_as_prompt_text(state.get("failed_steps", [])),
            plan_history=_as_prompt_text(plan_history),
            current_step_index=state.get("current_step_index", 0),
        )

        result = safe_invoke(
            create_plan_llm,
            messages=[SystemMessage(content=prompt)],
        )
        parsed = _extract_json_response(result)
        if isinstance(parsed, dict) and isinstance(parsed.get("steps"), list):
            parsed = parsed["steps"]
        if not isinstance(parsed, list) or not parsed:
            raise ValueError("create_plan must return a non-empty JSON list of plan steps.")

        required_fields = {
            "step_id", "objective", "tool_category", "focused_instruction",
            "required_inputs", "expected_result", "success_criteria",
            "fallback_action", "required",
        }
        valid_categories = set(TOOLS_BY_CATEGORY)
        valid_fallbacks = {"retry", "skip_optional", "replan", "fail"}
        plan: List[PlanStep] = []

        for raw_step in parsed:
            if not isinstance(raw_step, dict):
                raise ValueError("Every plan step must be a dictionary.")
            missing = required_fields - raw_step.keys()
            if missing:
                raise ValueError(f"Plan step is missing fields: {sorted(missing)}")
            raw_tool_category = raw_step["tool_category"]
            tool_category = TOOL_NAME_TO_CATEGORY.get(raw_tool_category,raw_tool_category,)
            if tool_category not in valid_categories:
                raise ValueError(f"Invalid tool category or tool name: {raw_tool_category}")
            if raw_step["fallback_action"] not in valid_fallbacks:
                raise ValueError(f"Invalid fallback action: {raw_step['fallback_action']}")

            step: PlanStep = {
                **raw_step,
                "tool_category": tool_category,
                "required_inputs": (
                    raw_step["required_inputs"]
                    if isinstance(raw_step["required_inputs"], dict)
                    else {}
                ),
                "required": bool(raw_step["required"]),
                "status": "pending",
                "attempts": 0,
            }
            plan.append(step)

        if plan_history:
            execution_counters["replans"] = execution_counters.get("replans", 0) + 1
            max_replans = state["execution_limits"].get("max_replans", 3)
            if execution_counters["replans"] > max_replans:
                raise RuntimeError("Maximum replanning count exceeded.")

        state["current_plan"] = plan
        state["plan_history"] = [*plan_history, [dict(step) for step in plan]]
        state["current_step_index"] = 0
        state["execution_counters"] = execution_counters
        state["next_action"] = "execute_plan"
        state["status"] = "running"
        state["failure_reason"] = None
        return state
    except Exception as exc:
        print(f"{RED}[NODE] [ERR]{RESET}", exc) if DEBUG else None
        traceback.print_exc() if DEBUG else None
        state["status"] = "failed"
        state["failure_reason"] = f"create_plan failed: {exc}"
        state["next_action"] = "__end__"
        return state


def execute_plan(state: AgentSchema) -> AgentSchema:
    """Ask the LLM for exactly one tool call for the active plan step."""
    print_function_name()
    _ensure_state_defaults(state)

    if state.get("status") == "failed":
        state["next_action"] = "__end__"
        return state

    limit_reason = _limit_failure(state)
    if limit_reason:
        state["status"] = "failed"
        state["failure_reason"] = limit_reason
        state["next_action"] = "__end__"
        return state

    plan = state["current_plan"]
    index = state["current_step_index"]
    if not plan or index < 0 or index >= len(plan):
        state["status"] = "failed"
        state["failure_reason"] = "No valid active plan step is available."
        state["next_action"] = "__end__"
        return state

    step = plan[index]
    max_attempts = state["execution_limits"].get("max_step_attempts", 3)
    if step.get("attempts", 0) >= max_attempts:
        _route_failed_step(
            state,
            step,
            f"Step {step.get('step_id')} reached the maximum attempt count.",
        )
        return state

    category = step.get("tool_category")
    selected_tool = TOOLS_BY_CATEGORY.get(category)
    if selected_tool is None:
        _route_failed_step(state, step, f"Unknown tool category: {category}")
        return state

    try:
        selected_tool = TOOLS_BY_CATEGORY.get(category)
        available_tool_name = selected_tool.name if selected_tool else "unknown"
        available_tool_description = selected_tool.description if selected_tool else ""

        prompt = prompts.EXECUTE_PLAN_PROMPT.format(
            current_step=_as_prompt_text(step),
            answer_requirements=_as_prompt_text(state.get("answer_requirements", {})),
            observations=_as_prompt_text(state.get("observations", [])),
            calculations=_as_prompt_text(state.get("calculations", [])),
            tool_call_log=_as_prompt_text(state.get("tool_call_log", [])),
            errors=_as_prompt_text(state.get("errors", [])),
            question=state.get("question", ""),
            attachment_path=state.get("attachment_path"),
            attachment_extension=state.get("attachment_extension"),
            attachment_supported=state.get("attachment_supported"),
            execution_counters=_as_prompt_text(state.get("execution_counters", {})),
            execution_limits=_as_prompt_text(state.get("execution_limits", {})),
            reviewer_feedback=state.get("reviewer_feedback") or "",
            available_tool_name=available_tool_name,
            available_tool_description=available_tool_description,
        )

        tool_enabled_llm = execute_plan_llm.bind_tools(
            [selected_tool],
            tool_choice=selected_tool.name,
            parallel_tool_calls=False,
        )
        result = safe_invoke(
            tool_enabled_llm,
            messages=[SystemMessage(content=prompt)],
        )
        if not isinstance(result, AIMessage):
            raise TypeError("execute_plan_llm did not return an AIMessage.")

        tool_calls = _get_tool_calls(result)
        if len(tool_calls) != 1:
            step["attempts"] = step.get("attempts", 0) + 1
            reason = f"Expected exactly one tool call, received {len(tool_calls)}."
            state["errors"].append({
                "node": "execute_plan",
                "error": reason,
                "timestamp": time.time(),
            })
            _route_failed_step(state, step, reason)
            return state

        state["messages"].append(result)
        state["next_action"] = "execute_plan"
        return state
    except Exception as exc:
        print(f"{RED}[NODE] [ERR]{RESET}", exc) if DEBUG else None
        traceback.print_exc() if DEBUG else None
        step["attempts"] = step.get("attempts", 0) + 1
        state["errors"].append({
            "node": "execute_plan",
            "error": str(exc),
            "traceback": traceback.format_exc(),
            "timestamp": time.time(),
        })
        _route_failed_step(state, step, f"Tool-call generation failed: {exc}")
        return state


def synthesize_candidate(state: AgentSchema) -> AgentSchema:
    """Synthesize a supported candidate answer from recorded evidence."""
    print_function_name()
    _ensure_state_defaults(state)
    if state.get("status") == "failed":
        return state

    try:
        prompt = prompts.SYNTHESIZE_CANDIDATE_PROMPT.format(
            question=state.get("question", ""),
            answer_requirements=_as_prompt_text(state.get("answer_requirements", {})),
            observations=_as_prompt_text(state.get("observations", [])),
            calculations=_as_prompt_text(state.get("calculations", [])),
            reviewer_feedback=state.get("reviewer_feedback") or "",
        )
        result = safe_invoke(
            synthesize_candidate_llm,
            messages=[SystemMessage(content=prompt)],
        )
        parsed = _extract_json_response(result)
        if not isinstance(parsed, dict):
            raise ValueError("synthesize_candidate must return a JSON object.")

        candidate_answer = parsed.get("candidate_answer")
        answer_support = parsed.get("answer_support", [])
        unresolved = parsed.get("unresolved_requirements", [])
        unsupported = parsed.get("unsupported_components", [])

        if not isinstance(candidate_answer, str) or not candidate_answer.strip():
            raise ValueError("candidate_answer is missing or empty.")
        if not isinstance(answer_support, list):
            raise ValueError("answer_support must be a list.")
        if not isinstance(unresolved, list) or not isinstance(unsupported, list):
            raise ValueError("Unresolved and unsupported fields must be lists.")

        validated_support: List[AnswerSupportEntry] = []
        for entry in answer_support:
            if not isinstance(entry, dict):
                continue
            validated_support.append({
                "answer_part": str(entry.get("answer_part", "")),
                "supporting_observations": [
                    str(value) for value in entry.get("supporting_observations", [])
                ] if isinstance(entry.get("supporting_observations", []), list) else [],
                "supporting_calculations": [
                    str(value) for value in entry.get("supporting_calculations", [])
                ] if isinstance(entry.get("supporting_calculations", []), list) else [],
            })

        state["candidate_answer"] = candidate_answer.strip()
        state["answer_support"] = validated_support
        state["unresolved_requirements"] = [str(value) for value in unresolved]
        state["unsupported_components"] = [str(value) for value in unsupported]
        return state
    except Exception as exc:
        print(f"{RED}[NODE] [ERR]{RESET}", exc) if DEBUG else None
        traceback.print_exc() if DEBUG else None
        state["status"] = "failed"
        state["failure_reason"] = f"synthesize_candidate failed: {exc}"
        return state


def review_candidate(state: AgentSchema) -> AgentSchema:
    """Review the candidate and return one supported routing decision."""
    print_function_name()
    _ensure_state_defaults(state)
    if state.get("status") == "failed":
        return state

    try:
        uncertainty_notes = [
            {
                "observation_id": observation.get("observation_id"),
                "uncertainty": observation.get("uncertainty"),
                "content": str(observation.get("content", ""))[:300],
            }
            for observation in state.get("observations", [])
            if observation.get("uncertainty") not in {None, "", 0, 0.0, "low", "none"}
        ]

        prompt = prompts.REVIEW_CANDIDATE_PROMPT.format(
            candidate_answer=state.get("candidate_answer", ""),
            answer_requirements=_as_prompt_text(state.get("answer_requirements", {})),
            observations=_as_prompt_text(state.get("observations", [])),
            calculations=_as_prompt_text(state.get("calculations", [])),
            tool_call_log=_as_prompt_text(state.get("tool_call_log", [])),
            unresolved_requirements=_as_prompt_text(state.get("unresolved_requirements", [])),
            unsupported_components=_as_prompt_text(state.get("unsupported_components", [])),
            question=state.get("question", ""),
            parser_uncertainty_notes=_as_prompt_text(uncertainty_notes),
            reviewer_feedback=state.get("reviewer_feedback") or "",
            answer_support=_as_prompt_text(state.get("answer_support", [])),
        )
        result = safe_invoke(
            review_candidate_llm,
            messages=[SystemMessage(content=prompt)],
        )
        parsed = _extract_json_response(result)
        if not isinstance(parsed, dict):
            raise ValueError("review_candidate must return a JSON object.")

        decision = parsed.get("reviewer_decision")
        feedback = str(parsed.get("reviewer_feedback", "")).strip()
        allowed = {
            "approve",
            "revise_answer",
            "recalculate",
            "gather_more_evidence",
            "revise_plan",
        }
        if decision not in allowed:
            raise ValueError(f"Invalid reviewer decision: {decision}")
        if decision != "approve" and not feedback:
            raise ValueError("Non-approved reviewer decisions require feedback.")

        if decision == "approve" and (
            state.get("unresolved_requirements")
            or state.get("unsupported_components")
        ):
            decision = "gather_more_evidence"
            feedback = (
                feedback
                or "The candidate contains unresolved requirements or unsupported components."
            )

        state["reviewer_decision"] = decision
        state["reviewer_feedback"] = feedback

        if decision in {"recalculate", "gather_more_evidence"} and state.get("current_plan"):
            index = min(
                state.get("current_step_index", 0),
                len(state["current_plan"]) - 1,
            )
            state["current_step_index"] = index
            state["current_plan"][index]["status"] = "pending"

        return state
    except Exception as exc:
        print(f"{RED}[NODE] [ERR]{RESET}", exc) if DEBUG else None
        traceback.print_exc() if DEBUG else None
        state["status"] = "failed"
        state["failure_reason"] = f"review_candidate failed: {exc}"
        return state


def format_output(state: AgentSchema) -> AgentSchema:
    """Format the approved candidate without changing its semantic meaning."""
    print_function_name()
    _ensure_state_defaults(state)

    if state.get("reviewer_decision") != "approve":
        state["status"] = "failed"
        state["failure_reason"] = "format_output called without reviewer approval."
        return state

    requirements = state.get("answer_requirements", {})
    previous_failure = ""

    for attempt in range(3):
        try:
            prompt = prompts.FORMAT_OUTPUT_PROMPT.format(
                candidate_answer=state.get("candidate_answer", ""),
                answer_requirements=_as_prompt_text(requirements),
                separator=requirements.get("separator"),
                expected_item_count=requirements.get("expected_item_count"),
                ordering=requirements.get("ordering"),
                units=requirements.get("units"),
                rounding=requirements.get("rounding"),
                capitalization=requirements.get("capitalization"),
                date_format=requirements.get("date_format"),
                required_prefix=requirements.get("required_prefix"),
                required_suffix=requirements.get("required_suffix"),
                prohibited_characters=_as_prompt_text(requirements.get("prohibited_characters", [])),
                prohibited_extra_text=_as_prompt_text(requirements.get("prohibited_extra_text", [])),
            )
            messages: List[BaseMessage] = [SystemMessage(content=prompt)]
            if previous_failure:
                messages.append(HumanMessage(content=f"Correct these formatting-check failures: {previous_failure}"))

            result = safe_invoke(format_output_llm, messages=messages)
            parsed = _extract_json_response(result)
            if not isinstance(parsed, dict):
                raise ValueError("format_output must return a JSON object.")

            summary = str(parsed.get("format_check_summary", "")).strip()
            final_output = parsed.get("final_output")
            if not isinstance(final_output, str) or not final_output.strip():
                raise ValueError("final_output is empty.")
            final_output = final_output.strip()

            failures: List[str] = []
            if "```" in final_output:
                failures.append("Markdown code fence is present")

            expected_count = requirements.get("expected_item_count")
            separator = requirements.get("separator")
            if isinstance(expected_count, int) and expected_count > 1 and separator:
                item_count = len([item for item in final_output.split(separator) if item.strip()])
                if item_count != expected_count:
                    failures.append(f"expected {expected_count} items but found {item_count}")

            prefix = requirements.get("required_prefix")
            suffix = requirements.get("required_suffix")
            units = requirements.get("units")
            if prefix and not final_output.startswith(prefix):
                failures.append("required prefix is missing")
            if suffix and not final_output.endswith(suffix):
                failures.append("required suffix is missing")
            if units and str(units) not in final_output:
                failures.append("required units are missing")

            for character in requirements.get("prohibited_characters", []):
                if character and character in final_output:
                    failures.append(f"prohibited character is present: {character!r}")

            lowered_output = final_output.lower()
            for prohibited_text in requirements.get("prohibited_extra_text", []):
                if prohibited_text and str(prohibited_text).lower() in lowered_output:
                    failures.append(f"prohibited text is present: {prohibited_text!r}")

            if failures:
                previous_failure = "; ".join(failures)
                continue

            state["format_check_summary"] = summary
            state["final_output"] = final_output
            state["status"] = "success"
            state["failure_reason"] = None
            return state
        except Exception as exc:
            previous_failure = str(exc)
            if attempt == 2:
                state["status"] = "failed"
                state["failure_reason"] = f"format_output failed: {exc}"
                return state

    state["status"] = "failed"
    state["failure_reason"] = f"format_output failed: {previous_failure}"
    return state


# Conditional Functions
def from_execute_plan_with_tools_to(
    state: AgentSchema,
) -> Literal[
    "execute_plan_tools_execution",
    "execute_plan",
    "create_plan",
    "synthesize_candidate",
    "__end__",
]:
    """Route a generated tool call to its handler or follow next_action."""
    if state.get("status") == "failed":
        return "__end__"

    messages = state.get("messages", [])
    if messages and isinstance(messages[-1], AIMessage) and _get_tool_calls(messages[-1]):
        return "execute_plan_tools_execution"
    return state.get("next_action", "__end__")


def from_execute_plan_tools_execution_to(
    state: AgentSchema,
) -> Literal["execute_plan", "create_plan", "synthesize_candidate", "__end__"]:
    """Route after tool execution using the handler's next_action."""
    return state.get("next_action", "__end__")


def from_review_candidate_to(
    state: AgentSchema,
) -> Literal[
    "format_output",
    "synthesize_candidate",
    "execute_plan",
    "create_plan",
    "__end__",
]:
    """Route according to reviewer_decision or end on failure."""
    if state.get("status") == "failed":
        return "__end__"

    mapping = {
        "approve": "format_output",
        "revise_answer": "synthesize_candidate",
        "recalculate": "execute_plan",
        "gather_more_evidence": "execute_plan",
        "revise_plan": "create_plan",
    }
    return mapping.get(state.get("reviewer_decision"), "__end__")


# Tool Handlers
def execute_plan_tools_execution(state: AgentSchema) -> AgentSchema:
    """Execute one pending tool call and update evidence and step state."""
    print_function_name()
    _ensure_state_defaults(state)

    plan = state["current_plan"]
    index = state["current_step_index"]
    if not plan or index < 0 or index >= len(plan):
        state["status"] = "failed"
        state["failure_reason"] = "Tool handler has no valid active plan step."
        state["next_action"] = "__end__"
        return state

    step = plan[index]
    messages = state["messages"]
    if not messages or not isinstance(messages[-1], AIMessage):
        _route_failed_step(state, step, "Tool handler did not receive an AI tool-call message.")
        return state

    calls = _get_tool_calls(messages[-1])
    if len(calls) != 1:
        step["attempts"] = step.get("attempts", 0) + 1
        _route_failed_step(
            state,
            step,
            f"Tool handler expected exactly one call, received {len(calls)}.",
        )
        return state

    start_time = time.time()
    tool_name = "unknown"
    arguments: Dict[str, Any] = {}
    call_id = f"call_{uuid.uuid4().hex}"

    try:
        tool_name, arguments, call_id = _normalise_tool_call(calls[0])
        expected_name = TOOL_CATEGORY_TO_NAME.get(step.get("tool_category"))
        if tool_name != expected_name:
            raise ValueError(
                f"Tool '{tool_name}' does not match planned category "
                f"'{step.get('tool_category')}' (expected '{expected_name}')."
            )

        selected_tool = TOOLS_BY_NAME.get(tool_name)
        if selected_tool is None:
            raise ValueError(f"Unregistered tool requested: {tool_name}")

        _register_attachment_path(state.get("attachment_path"))
        step["attempts"] = step.get("attempts", 0) + 1

        previous_logs = list(state["tool_call_log"])
        result = selected_tool.invoke(arguments)
        end_time = time.time()

        if isinstance(result, dict) and result.get("error"):
            raise RuntimeError(str(result["error"]))

        state["messages"].append(ToolMessage(
            content=json.dumps(result, ensure_ascii=False, default=_json_safe),
            tool_call_id=call_id,
            name=tool_name,
        ))

        state["tool_call_log"].append({
            "tool_name": tool_name,
            "arguments": arguments,
            "start_time": start_time,
            "end_time": end_time,
            "status": "success",
            "error": None,
        })

        counters = state["execution_counters"]
        counters["total_tool_calls"] = counters.get("total_tool_calls", 0) + 1
        counters[tool_name] = counters.get(tool_name, 0) + 1

        if tool_name == "web_search_tool":
            query = arguments.get("query")
            if any(
                log.get("tool_name") == tool_name
                and log.get("arguments", {}).get("query") == query
                for log in previous_logs[-5:]
            ):
                counters["repeated_searches"] = counters.get("repeated_searches", 0) + 1

        action_key = json.dumps(
            {"tool": tool_name, "arguments": arguments},
            sort_keys=True,
            default=str,
        )
        if any(
            json.dumps(
                {"tool": log.get("tool_name"), "arguments": log.get("arguments", {})},
                sort_keys=True,
                default=str,
            ) == action_key
            for log in previous_logs[-10:]
        ):
            counters["duplicate_actions"] = counters.get("duplicate_actions", 0) + 1

        observation_base = f"obs_{uuid.uuid4().hex}"
        if tool_name == "web_search_tool":
            for result_index, item in enumerate(result):
                state["observations"].append({
                    "observation_id": f"{observation_base}_{result_index}",
                    "content": str(item.get("snippet", "")),
                    "value": item.get("title"),
                    "source": str(item.get("url", "")),
                    "source_location": f"search result {result_index + 1}",
                    "tool_name": tool_name,
                    "confidence": 0.0,
                    "uncertainty": "Unverified search-result snippet",
                })

        elif tool_name == "webpage_read_tool":
            for passage_index, passage in enumerate(result.get("passages", [])):
                state["observations"].append({
                    "observation_id": f"{observation_base}_{passage_index}",
                    "content": str(passage.get("content", "")),
                    "value": None,
                    "source": str(result.get("url", arguments.get("url", ""))),
                    "source_location": str(passage.get("source_location", "unknown")),
                    "tool_name": tool_name,
                    "confidence": 1.0,
                    "uncertainty": "Deterministically extracted passage",
                })
            if result.get("downloaded_path"):
                state["observations"].append({
                    "observation_id": f"{observation_base}_download",
                    "content": "Downloaded a supported file to the managed temporary directory.",
                    "value": result["downloaded_path"],
                    "source": str(result.get("url", arguments.get("url", ""))),
                    "source_location": "managed temporary file",
                    "tool_name": tool_name,
                    "confidence": 1.0,
                    "uncertainty": "File transfer completed; content not yet inspected",
                })

        elif tool_name == "deterministic_parser_tool":
            for item_index, item in enumerate(result.get("items", [])):
                state["observations"].append({
                    "observation_id": f"{observation_base}_{item_index}",
                    "content": json.dumps(item.get("content"), ensure_ascii=False, default=str),
                    "value": item.get("content"),
                    "source": str(result.get("file_path", arguments.get("file_path", ""))),
                    "source_location": str(item.get("source_location", "unknown")),
                    "tool_name": tool_name,
                    "confidence": 1.0,
                    "uncertainty": "Deterministically parsed content",
                })

        elif tool_name == "multimodal_inspection_tool":
            for item_index, item in enumerate(result):
                state["observations"].append({
                    "observation_id": f"{observation_base}_{item_index}",
                    "content": str(item.get("content", "")),
                    "value": item.get("value"),
                    "source": str(arguments.get("file_path", "")),
                    "source_location": str(item.get("source_location", "unknown")),
                    "tool_name": tool_name,
                    "confidence": float(item.get("confidence", 0.5)),
                    "uncertainty": str(item.get("uncertainty", "unspecified")),
                })

        elif tool_name == "secure_zip_tool":
            state["observations"].append({
                "observation_id": observation_base,
                "content": (
                    "ZIP member listing"
                    if result.get("action") == "list"
                    else "ZIP member extracted"
                ),
                "value": result.get("members") if result.get("action") == "list" else result.get("extracted_path"),
                "source": str(arguments.get("zip_path", "")),
                "source_location": (
                    "archive root"
                    if result.get("action") == "list"
                    else str(arguments.get("member_path", "unknown member"))
                ),
                "tool_name": tool_name,
                "confidence": 1.0,
                "uncertainty": "Deterministic archive operation",
            })

        elif tool_name == "restricted_python_tool":
            stdout = str(result.get("stdout", "")).strip()
            try:
                calculation_result = json.loads(stdout) if stdout else None
            except json.JSONDecodeError:
                calculation_result = stdout
            state["calculations"].append({
                "calculation_id": f"calc_{uuid.uuid4().hex}",
                "description": step.get("objective", "Python calculation"),
                "inputs": step.get("required_inputs", {}),
                "method": str(arguments.get("code", "")),
                "result": calculation_result,
                "units": str(state.get("answer_requirements", {}).get("units") or ""),
                "completed": bool(result.get("completed")),
                "error": result.get("error"),
            })

        step["status"] = "completed"
        next_index = index + 1
        if next_index < len(plan):
            state["current_step_index"] = next_index
            state["next_action"] = "execute_plan"
        else:
            required_incomplete = [
                candidate
                for candidate in plan
                if candidate.get("required", True)
                and candidate.get("status") != "completed"
            ]
            if required_incomplete:
                raise RuntimeError("One or more required plan steps remain incomplete.")
            state["next_action"] = "synthesize_candidate"

        state["status"] = "running"
        state["failure_reason"] = None
        return state

    except Exception as exc:
        end_time = time.time()
        error_text = str(exc)
        state["tool_call_log"].append({
            "tool_name": tool_name,
            "arguments": arguments,
            "start_time": start_time,
            "end_time": end_time,
            "status": "error",
            "error": error_text,
        })
        state["errors"].append({
            "tool_name": tool_name,
            "arguments": arguments,
            "error": error_text,
            "traceback": traceback.format_exc(),
            "timestamp": end_time,
        })
        state["messages"].append(ToolMessage(
            content=json.dumps({"error": error_text}, ensure_ascii=False),
            tool_call_id=call_id,
            name=tool_name if tool_name != "unknown" else None,
        ))
        if step.get("attempts", 0) == 0:
            step["attempts"] = 1
        _route_failed_step(
            state,
            step,
            f"Step {step.get('step_id')} failed: {error_text}",
        )
        return state


# Graph
gaia_solver_graph = StateGraph(AgentSchema)

gaia_solver_graph.add_node("analyze_task", analyze_task)
gaia_solver_graph.add_node("create_plan", create_plan)
gaia_solver_graph.add_node("execute_plan", execute_plan)
gaia_solver_graph.add_node("execute_plan_tools_execution", execute_plan_tools_execution)
gaia_solver_graph.add_node("synthesize_candidate", synthesize_candidate)
gaia_solver_graph.add_node("review_candidate", review_candidate)
gaia_solver_graph.add_node("format_output", format_output)

gaia_solver_graph.add_edge(START, "analyze_task")
gaia_solver_graph.add_edge("analyze_task", "create_plan")
gaia_solver_graph.add_edge("create_plan", "execute_plan")
gaia_solver_graph.add_conditional_edges(
    "execute_plan",
    from_execute_plan_with_tools_to,
    {
        "execute_plan_tools_execution": "execute_plan_tools_execution",
        "execute_plan": "execute_plan",
        "create_plan": "create_plan",
        "synthesize_candidate": "synthesize_candidate",
        "__end__": "__end__",
    },
)
gaia_solver_graph.add_conditional_edges(
    "execute_plan_tools_execution",
    from_execute_plan_tools_execution_to,
    {
        "execute_plan": "execute_plan",
        "create_plan": "create_plan",
        "synthesize_candidate": "synthesize_candidate",
        "__end__": "__end__",
    },
)
gaia_solver_graph.add_edge("synthesize_candidate", "review_candidate")
gaia_solver_graph.add_conditional_edges(
    "review_candidate",
    from_review_candidate_to,
    {
        "format_output": "format_output",
        "synthesize_candidate": "synthesize_candidate",
        "execute_plan": "execute_plan",
        "create_plan": "create_plan",
        "__end__": "__end__",
    },
)
gaia_solver_graph.add_edge("format_output", END)

gaia_solver_app = gaia_solver_graph.compile()


# Testing
if __name__ == "__main__":
    from IPython.display import Image as GraphImage

    parent_dir = Path(__file__).resolve().parent
    graph_dir = parent_dir / "graphs"
    graph_dir.mkdir(parents=True, exist_ok=True)
    graph_png = gaia_solver_app.get_graph().draw_mermaid_png(max_retries=5, retry_delay=2.0)
    GraphImage(graph_png)
    (graph_dir / "gaia_solver_app.png").write_bytes(graph_png)

    from langsmith import Client

    os.environ["LANGCHAIN_PROJECT"] = "gaia_solver"
    os.environ["LANGSMITH_PROJECT"] = "gaia_solver"
    client = Client()

    config = {
        "recursion_limit": 100,
        "configurable": {
            "user_id": "gaia_solver",
            "run_name": "gaia_solver",
            "thread_id": "gaia_solver",
        },
    }

    initial_state: AgentSchema = {
        "question": "",  # Add a GAIA question.
        "attachment_path": None,
        "messages": [],
        "status": "running",
    }
    response = gaia_solver_app.invoke(initial_state, config=config)

    print(f"{BLUE}[MAIN] [INFO]{RESET} Response") if DEBUG else None
    if DEBUG:
        for key, value in response.items():
            print(f"    {key}: {value}")

    
    # pip install tavily-python requests beautifulsoup4 pypdf openpyxl python-docx python-pptx pymupdf